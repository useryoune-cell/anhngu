# -*- coding: utf-8 -*-
from pathlib import Path
from datetime import datetime
import json
import os
import random
import re
import threading
import urllib.error
import urllib.request as urllib_request
import uuid
from urllib.parse import parse_qs, urlparse
import sqlite3
import sys
import time

try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None
try:
    from docx import Document
except ImportError:
    Document = None

from flask import (
    Flask,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_from_directory,
    session,
    url_for,
)
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename


def configure_utf8_streams():
    for stream_name in ("stdout", "stderr"):
        stream = getattr(sys, stream_name, None)
        if stream and hasattr(stream, "reconfigure"):
            try:
                stream.reconfigure(encoding="utf-8")
            except (LookupError, ValueError):
                pass


configure_utf8_streams()

BASE_DIR = Path(__file__).resolve().parent
if load_dotenv:
    load_dotenv(BASE_DIR / ".env")

DATA_DIR = Path(os.environ.get("LINGO_DATA_DIR", BASE_DIR)).resolve()
STATIC_UPLOAD_DIR = BASE_DIR / "static" / "uploads"
UPLOAD_ROOT = Path(os.environ.get("LINGO_UPLOAD_ROOT", STATIC_UPLOAD_DIR)).resolve()
DB_PATH = Path(os.environ.get("LINGO_DB_PATH", DATA_DIR / "users.db")).resolve()
AVATAR_DIR = UPLOAD_ROOT / "avatars"
SKILL_FILE_DIR = UPLOAD_ROOT / "skills" / "files"
EXAM_FILE_DIR = UPLOAD_ROOT / "exams" / "rooms"
ALLOWED_AVATAR_EXTENSIONS = {"png", "jpg", "jpeg", "webp", "gif"}
ALLOWED_MATERIAL_EXTENSIONS = {
    "pdf",
    "png",
    "jpg",
    "jpeg",
    "webp",
    "gif",
    "mp4",
    "webm",
    "mov",
    "mp3",
    "wav",
    "doc",
    "docx",
    "ppt",
    "pptx",
}

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-secret-key")
app.config["JSON_AS_ASCII"] = False
app.json.ensure_ascii = False

RANK_KIND_QUESTION_SECONDS = 6
RANK_KIND_INTERMISSION_SECONDS = 1.6
RANK_KIND_WIN_POINTS = 10
RANK_KIND_LOCK = threading.Lock()
RANK_KIND_QUEUE = []
RANK_KIND_MATCHES = {}


def prepare_runtime_storage():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)
    if UPLOAD_ROOT == STATIC_UPLOAD_DIR.resolve():
        return
    STATIC_UPLOAD_DIR.parent.mkdir(parents=True, exist_ok=True)
    if STATIC_UPLOAD_DIR.exists() or STATIC_UPLOAD_DIR.is_symlink():
        return
    try:
        STATIC_UPLOAD_DIR.symlink_to(UPLOAD_ROOT, target_is_directory=True)
    except OSError:
        STATIC_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


prepare_runtime_storage()


@app.after_request
def force_utf8_response(response):
    utf8_mimetypes = {
        "application/javascript",
        "application/json",
        "text/css",
        "text/html",
        "text/javascript",
        "text/plain",
    }
    if response.mimetype in utf8_mimetypes:
        response.headers["Content-Type"] = f"{response.mimetype}; charset=utf-8"
    return response


def get_db():
    db = getattr(g, "_database", None)
    if db is None:
        db = g._database = sqlite3.connect(DB_PATH)
        db.row_factory = sqlite3.Row
    return db


def init_db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    AVATAR_DIR.mkdir(parents=True, exist_ok=True)
    SKILL_FILE_DIR.mkdir(parents=True, exist_ok=True)
    EXAM_FILE_DIR.mkdir(parents=True, exist_ok=True)
    db = get_db()
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fullname TEXT NOT NULL,
            username TEXT UNIQUE,
            email TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            role TEXT NOT NULL,
            class TEXT
        )
        """
    )
    existing_user_columns = {
        row["name"] for row in db.execute("PRAGMA table_info(users)").fetchall()
    }
    if "phone" not in existing_user_columns:
        db.execute("ALTER TABLE users ADD COLUMN phone TEXT")
    if "avatar_path" not in existing_user_columns:
        db.execute("ALTER TABLE users ADD COLUMN avatar_path TEXT")
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            room_name TEXT NOT NULL DEFAULT 'Phòng thi',
            score INTEGER NOT NULL DEFAULT 0,
            taken_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS quiz_packages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            difficulty TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS quiz_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            package_id INTEGER NOT NULL,
            position INTEGER NOT NULL,
            question_text TEXT NOT NULL,
            option_a TEXT NOT NULL,
            option_b TEXT NOT NULL,
            option_c TEXT NOT NULL,
            option_d TEXT NOT NULL,
            correct_answer TEXT NOT NULL,
            FOREIGN KEY (package_id) REFERENCES quiz_packages (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS quiz_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            package_id INTEGER NOT NULL,
            correct_count INTEGER NOT NULL,
            total_questions INTEGER NOT NULL,
            percentage INTEGER NOT NULL,
            bonus_points INTEGER NOT NULL DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id),
            FOREIGN KEY (package_id) REFERENCES quiz_packages (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS rank_kind_results (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            opponent_id INTEGER,
            package_id INTEGER,
            result TEXT NOT NULL,
            score INTEGER NOT NULL DEFAULT 0,
            opponent_score INTEGER NOT NULL DEFAULT 0,
            rank_points INTEGER NOT NULL DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id),
            FOREIGN KEY (opponent_id) REFERENCES users (id),
            FOREIGN KEY (package_id) REFERENCES quiz_packages (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS listening_lessons (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            passage TEXT NOT NULL,
            quiz_question TEXT NOT NULL,
            option_a TEXT NOT NULL,
            option_b TEXT NOT NULL,
            option_c TEXT NOT NULL,
            option_d TEXT NOT NULL,
            correct_answer TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS listening_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            lesson_id INTEGER NOT NULL,
            selected_answer TEXT,
            score INTEGER NOT NULL DEFAULT 0,
            completed_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(user_id, lesson_id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            FOREIGN KEY (lesson_id) REFERENCES listening_lessons (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS grammar_lessons (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            content TEXT NOT NULL,
            material_type TEXT NOT NULL,
            material_url TEXT,
            material_path TEXT,
            material_name TEXT,
            quiz_question TEXT NOT NULL,
            option_a TEXT NOT NULL,
            option_b TEXT NOT NULL,
            option_c TEXT NOT NULL,
            option_d TEXT NOT NULL,
            correct_answer TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS grammar_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            lesson_id INTEGER NOT NULL,
            selected_answer TEXT,
            score INTEGER NOT NULL DEFAULT 0,
            completed_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(user_id, lesson_id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            FOREIGN KEY (lesson_id) REFERENCES grammar_lessons (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS writing_tasks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            prompt TEXT NOT NULL,
            time_limit_minutes INTEGER NOT NULL DEFAULT 30,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS writing_submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            task_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            content TEXT NOT NULL,
            ai_score INTEGER,
            ai_feedback TEXT,
            teacher_score INTEGER,
            teacher_feedback TEXT,
            status TEXT NOT NULL DEFAULT 'waiting_teacher',
            submitted_at TEXT DEFAULT CURRENT_TIMESTAMP,
            graded_at TEXT,
            UNIQUE(user_id, task_id),
            FOREIGN KEY (task_id) REFERENCES writing_tasks (id),
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS speaking_lessons (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            passage TEXT NOT NULL,
            word_notes_json TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS speaking_attempts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            lesson_id INTEGER NOT NULL,
            transcript TEXT,
            score INTEGER NOT NULL DEFAULT 0,
            missed_words_json TEXT,
            completed_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(user_id, lesson_id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            FOREIGN KEY (lesson_id) REFERENCES speaking_lessons (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_speaking_topics (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            topic_prompt TEXT NOT NULL,
            level TEXT NOT NULL DEFAULT 'A2-B1',
            opening_question TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_speaking_sessions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            topic_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            conversation_json TEXT,
            score INTEGER,
            feedback TEXT,
            pronunciation_feedback TEXT,
            unclear_words_json TEXT,
            status TEXT NOT NULL DEFAULT 'completed',
            started_at TEXT DEFAULT CURRENT_TIMESTAMP,
            completed_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (topic_id) REFERENCES exam_speaking_topics (id),
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_rooms (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            teacher_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            description TEXT,
            start_at TEXT NOT NULL,
            duration_minutes INTEGER NOT NULL DEFAULT 45,
            source_file_name TEXT,
            total_questions INTEGER NOT NULL DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (teacher_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_room_questions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            room_id INTEGER NOT NULL,
            position INTEGER NOT NULL,
            question_type TEXT NOT NULL,
            passage TEXT,
            question_text TEXT NOT NULL,
            option_a TEXT NOT NULL,
            option_b TEXT NOT NULL,
            option_c TEXT NOT NULL,
            option_d TEXT NOT NULL,
            correct_answer TEXT NOT NULL,
            FOREIGN KEY (room_id) REFERENCES exam_rooms (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_room_registrations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            room_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            registered_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(room_id, user_id),
            FOREIGN KEY (room_id) REFERENCES exam_rooms (id),
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS exam_room_submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            room_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            answers_json TEXT,
            correct_count INTEGER NOT NULL DEFAULT 0,
            total_questions INTEGER NOT NULL DEFAULT 0,
            score REAL NOT NULL DEFAULT 0,
            submitted_at TEXT DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(room_id, user_id),
            FOREIGN KEY (room_id) REFERENCES exam_rooms (id),
            FOREIGN KEY (user_id) REFERENCES users (id)
        )
        """
    )
    teacher_exists = db.execute(
        "SELECT id FROM users WHERE username = ? OR email = ?",
        ("giaovien01", "giaovien@lingo.vn"),
    ).fetchone()
    if not teacher_exists:
        db.execute(
            "INSERT INTO users (fullname, username, email, password_hash, role, class) VALUES (?, ?, ?, ?, ?, ?)",
            (
                "Giáo Viên Demo",
                "giaovien01",
                "giaovien@lingo.vn",
                generate_password_hash("123456"),
                "teacher",
                None,
            ),
        )
    db.commit()


def get_current_user():
    init_db()
    user_id = session.get("user_id")
    if not user_id:
        return None

    db = get_db()
    user = db.execute(
        "SELECT id, fullname, email, username, role, class, phone, avatar_path FROM users WHERE id = ?",
        (user_id,),
    ).fetchone()
    if not user:
        session.clear()
        return None
    return user


def avatar_is_allowed(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_AVATAR_EXTENSIONS


def material_file_is_allowed(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_MATERIAL_EXTENSIONS


def extract_youtube_id(raw_url):
    parsed = urlparse(raw_url.strip())
    hostname = parsed.netloc.lower().replace("www.", "")

    if hostname == "youtu.be":
        return parsed.path.strip("/").split("/")[0]

    if hostname in {"youtube.com", "m.youtube.com", "music.youtube.com"}:
        query_id = parse_qs(parsed.query).get("v", [""])[0]
        if query_id:
            return query_id
        if parsed.path.startswith("/embed/") or parsed.path.startswith("/shorts/"):
            parts = [part for part in parsed.path.split("/") if part]
            if len(parts) >= 2:
                return parts[1]

    return ""


def normalize_embed_url(raw_url):
    url = raw_url.strip()
    if not url.startswith(("http://", "https://")):
        url = f"https://{url}"

    parsed = urlparse(url)
    hostname = parsed.netloc.lower().replace("www.", "")
    if hostname == "drive.google.com":
        path_parts = [part for part in parsed.path.split("/") if part]
        if len(path_parts) >= 3 and path_parts[0] == "file" and path_parts[1] == "d":
            return f"https://drive.google.com/file/d/{path_parts[2]}/preview"
        file_id = parse_qs(parsed.query).get("id", [""])[0]
        if file_id:
            return f"https://drive.google.com/file/d/{file_id}/preview"

    return url


def get_quiz_package_or_404(package_id):
    return get_db().execute(
        """
        SELECT quiz_packages.*, users.fullname AS teacher_name
        FROM quiz_packages
        JOIN users ON users.id = quiz_packages.teacher_id
        WHERE quiz_packages.id = ?
        """,
        (package_id,),
    ).fetchone()


def get_quiz_questions(package_id):
    return get_db().execute(
        """
        SELECT *
        FROM quiz_questions
        WHERE package_id = ?
        ORDER BY position ASC, id ASC
        """,
        (package_id,),
    ).fetchall()


def get_all_quiz_packages(user=None):
    packages = get_db().execute(
        """
        SELECT quiz_packages.*, users.fullname AS teacher_name,
               COUNT(quiz_questions.id) AS question_count
        FROM quiz_packages
        JOIN users ON users.id = quiz_packages.teacher_id
        LEFT JOIN quiz_questions ON quiz_questions.package_id = quiz_packages.id
        GROUP BY quiz_packages.id
        ORDER BY quiz_packages.updated_at DESC, quiz_packages.id DESC
        """
    ).fetchall()

    if not user or user["role"] != "student":
        return packages

    best_rows = get_db().execute(
        """
        SELECT package_id, MAX(percentage) AS best_percentage, SUM(bonus_points) AS total_bonus
        FROM quiz_attempts
        WHERE user_id = ?
        GROUP BY package_id
        """,
        (user["id"],),
    ).fetchall()
    stats = {row["package_id"]: row for row in best_rows}
    result = []
    for row in packages:
        item = dict(row)
        stat = stats.get(row["id"])
        item["best_percentage"] = stat["best_percentage"] if stat else None
        item["total_bonus"] = stat["total_bonus"] if stat else 0
        result.append(item)
    return result


def get_quiz_form_questions():
    questions = []
    for index in range(1, 11):
        question_text = request.form.get(f"question_{index}", "").strip()
        option_a = request.form.get(f"option_{index}_a", "").strip()
        option_b = request.form.get(f"option_{index}_b", "").strip()
        option_c = request.form.get(f"option_{index}_c", "").strip()
        option_d = request.form.get(f"option_{index}_d", "").strip()
        correct_answer = request.form.get(f"correct_{index}", "").strip().upper()
        if not all([question_text, option_a, option_b, option_c, option_d]) or correct_answer not in {"A", "B", "C", "D"}:
            return None
        questions.append(
            {
                "position": index,
                "question_text": question_text,
                "option_a": option_a,
                "option_b": option_b,
                "option_c": option_c,
                "option_d": option_d,
                "correct_answer": correct_answer,
            }
        )
    return questions


def save_quiz_questions(package_id, questions):
    db = get_db()
    db.execute("DELETE FROM quiz_questions WHERE package_id = ?", (package_id,))
    for question in questions:
        db.execute(
            """
            INSERT INTO quiz_questions
                (package_id, position, question_text, option_a, option_b, option_c, option_d, correct_answer)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                package_id,
                question["position"],
                question["question_text"],
                question["option_a"],
                question["option_b"],
                question["option_c"],
                question["option_d"],
                question["correct_answer"],
            ),
        )


def get_rank_kind_packages(user=None):
    return [dict(package) for package in get_all_quiz_packages(user) if package["question_count"] == 10]


def get_rank_kind_player(user):
    display_name = user["fullname"] or user["username"] or user["email"] or "Người chơi"
    return {
        "id": str(user["id"]),
        "user_id": user["id"],
        "displayName": display_name,
        "avatarPath": user["avatar_path"] or "",
    }


def build_rank_kind_questions(package_id):
    rows = list(get_quiz_questions(package_id))
    if len(rows) != 10:
        return []
    questions = []
    for row in rows:
        questions.append(
            {
                "id": str(row["id"]),
                "prompt": row["question_text"],
                "answers": [
                    {"key": "A", "text": row["option_a"]},
                    {"key": "B", "text": row["option_b"]},
                    {"key": "C", "text": row["option_c"]},
                    {"key": "D", "text": row["option_d"]},
                ],
                "correctKey": row["correct_answer"],
            }
        )
    random.shuffle(questions)
    return questions


def get_rank_kind_leaderboard(limit=20):
    rows = get_db().execute(
        """
        SELECT
            users.id AS user_id,
            users.fullname,
            users.username,
            users.email,
            users.avatar_path,
            COALESCE(SUM(rank_kind_results.rank_points), 0) AS rank_points,
            SUM(CASE WHEN rank_kind_results.result = 'win' THEN 1 ELSE 0 END) AS wins,
            SUM(CASE WHEN rank_kind_results.result = 'lose' THEN 1 ELSE 0 END) AS losses,
            SUM(CASE WHEN rank_kind_results.result = 'draw' THEN 1 ELSE 0 END) AS draws,
            COUNT(rank_kind_results.id) AS matches
        FROM rank_kind_results
        JOIN users ON users.id = rank_kind_results.user_id
        GROUP BY users.id
        ORDER BY rank_points DESC, wins DESC, matches ASC, users.fullname ASC
        LIMIT ?
        """,
        (limit,),
    ).fetchall()
    leaderboard = []
    for index, row in enumerate(rows, start=1):
        leaderboard.append(
            {
                "rank": index,
                "userId": row["user_id"],
                "displayName": row["fullname"] or row["username"] or row["email"] or "Người chơi",
                "avatarPath": row["avatar_path"] or "",
                "rankPoints": row["rank_points"] or 0,
                "wins": row["wins"] or 0,
                "losses": row["losses"] or 0,
                "draws": row["draws"] or 0,
                "matches": row["matches"] or 0,
            }
        )
    return leaderboard


def get_rank_kind_package_or_first(package_id, user=None):
    packages = get_rank_kind_packages(user)
    if package_id:
        for package in packages:
            if package["id"] == package_id:
                return package
    return packages[0] if packages else None


def rank_kind_public_player(player, score=0, ready=False):
    return {
        **player,
        "score": score,
        "ready": ready,
        "avatarUrl": url_for("static", filename=player["avatarPath"]) if player.get("avatarPath") else "",
    }


def record_rank_kind_result(match):
    if match.get("recorded"):
        return
    match["recorded"] = True
    players = match["players"]
    scores = match["scores"]
    winner_id = match.get("winnerId")
    db = get_db()
    for player in players:
        opponent = next((item for item in players if item["id"] != player["id"]), None)
        score = scores.get(player["id"], 0)
        opponent_score = scores.get(opponent["id"], 0) if opponent else 0
        result = "draw"
        rank_points = 0
        if winner_id:
            result = "win" if winner_id == player["id"] else "lose"
            rank_points = RANK_KIND_WIN_POINTS if result == "win" else 0
        db.execute(
            """
            INSERT INTO rank_kind_results
                (user_id, opponent_id, package_id, result, score, opponent_score, rank_points)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                player["user_id"],
                opponent["user_id"] if opponent else None,
                match["packageId"],
                result,
                score,
                opponent_score,
                rank_points,
            ),
        )
    db.commit()


def resolve_rank_kind_question(match, winner_id=None, selected_by=None, selected_key=None, reason="timeout"):
    if match["current"].get("resolved"):
        return
    match["current"]["resolved"] = True
    match["current"]["winnerId"] = winner_id
    match["current"]["selectedBy"] = selected_by
    match["current"]["selectedKey"] = selected_key
    match["current"]["reason"] = reason
    match["current"]["resultUntil"] = time.time() + RANK_KIND_INTERMISSION_SECONDS
    if winner_id:
        match["scores"][winner_id] = match["scores"].get(winner_id, 0) + 1


def start_rank_kind_question(match, index):
    question = match["questions"][index]
    match["questionIndex"] = index
    match["current"] = {
        "id": question["id"],
        "deadline": time.time() + RANK_KIND_QUESTION_SECONDS,
        "answered": set(),
        "resolved": False,
        "winnerId": None,
        "selectedBy": None,
        "selectedKey": None,
        "reason": "",
        "resultUntil": 0,
    }


def finish_rank_kind_match(match):
    if match["state"] == "finished":
        return
    left, right = match["players"]
    left_score = match["scores"].get(left["id"], 0)
    right_score = match["scores"].get(right["id"], 0)
    winner_id = None
    if left_score > right_score:
        winner_id = left["id"]
    elif right_score > left_score:
        winner_id = right["id"]
    match["winnerId"] = winner_id
    match["state"] = "finished"
    match["finishedAt"] = time.time()
    record_rank_kind_result(match)


def advance_rank_kind_match(match):
    if match["state"] != "playing":
        return
    current = match.get("current") or {}
    now = time.time()
    if not current.get("resolved") and now >= current.get("deadline", 0):
        resolve_rank_kind_question(match, reason="timeout")
        return
    if current.get("resolved") and now >= current.get("resultUntil", 0):
        next_index = match["questionIndex"] + 1
        if next_index >= len(match["questions"]):
            finish_rank_kind_match(match)
        else:
            start_rank_kind_question(match, next_index)


def serialize_rank_kind_match(match, user_id):
    advance_rank_kind_match(match)
    scores = match["scores"]
    players = [
        rank_kind_public_player(player, scores.get(player["id"], 0), player["id"] in match["ready"])
        for player in match["players"]
    ]
    payload = {
        "status": match["state"],
        "matchId": match["id"],
        "packageTitle": match["packageTitle"],
        "players": players,
        "scores": scores,
        "questionNumber": match["questionIndex"] + 1 if match["questionIndex"] >= 0 else 0,
        "totalQuestions": len(match["questions"]),
        "localPlayerId": str(user_id),
    }
    if match["state"] == "playing":
        question = match["questions"][match["questionIndex"]]
        current = match["current"]
        payload["question"] = {
            "id": question["id"],
            "prompt": question["prompt"],
            "answers": question["answers"],
            "deadline": current["deadline"],
            "serverNow": time.time(),
            "resolved": current["resolved"],
        }
        if current["resolved"]:
            payload["question"].update(
                {
                    "correctKey": question["correctKey"],
                    "winnerId": current["winnerId"],
                    "selectedBy": current["selectedBy"],
                    "selectedKey": current["selectedKey"],
                    "reason": current["reason"],
                    "nextInMs": max(0, int((current["resultUntil"] - time.time()) * 1000)),
                }
            )
    if match["state"] == "finished":
        payload["winnerId"] = match.get("winnerId")
        payload["leaderboard"] = get_rank_kind_leaderboard()
    return payload


def create_rank_kind_match(first_ticket, second_player, package):
    questions = build_rank_kind_questions(package["id"])
    match_id = uuid.uuid4().hex
    players = [first_ticket["player"], second_player]
    match = {
        "id": match_id,
        "packageId": package["id"],
        "packageTitle": package["title"],
        "state": "ready",
        "players": players,
        "ready": set(),
        "scores": {players[0]["id"]: 0, players[1]["id"]: 0},
        "questions": questions,
        "questionIndex": -1,
        "current": None,
        "winnerId": None,
        "recorded": False,
        "createdAt": time.time(),
    }
    RANK_KIND_MATCHES[match_id] = match
    return match


def get_learning_attempt(table_name, user_id, lesson_id):
    return get_db().execute(
        f"SELECT * FROM {table_name} WHERE user_id = ? AND lesson_id = ?",
        (user_id, lesson_id),
    ).fetchone()


def get_listening_lesson_or_404(lesson_id):
    return get_db().execute(
        """
        SELECT listening_lessons.*, users.fullname AS teacher_name
        FROM listening_lessons
        JOIN users ON users.id = listening_lessons.teacher_id
        WHERE listening_lessons.id = ?
        """,
        (lesson_id,),
    ).fetchone()


def get_all_listening_lessons(user=None):
    rows = get_db().execute(
        """
        SELECT listening_lessons.*, users.fullname AS teacher_name
        FROM listening_lessons
        JOIN users ON users.id = listening_lessons.teacher_id
        ORDER BY listening_lessons.updated_at DESC, listening_lessons.id DESC
        """
    ).fetchall()
    if not user or user["role"] != "student":
        return rows

    completed_ids = {
        row["lesson_id"]
        for row in get_db().execute(
            "SELECT lesson_id FROM listening_attempts WHERE user_id = ?",
            (user["id"],),
        ).fetchall()
    }
    lessons = []
    for row in rows:
        lesson = dict(row)
        lesson["is_completed"] = row["id"] in completed_ids
        lessons.append(lesson)
    return lessons


def get_grammar_lesson_or_404(lesson_id):
    return get_db().execute(
        """
        SELECT grammar_lessons.*, users.fullname AS teacher_name
        FROM grammar_lessons
        JOIN users ON users.id = grammar_lessons.teacher_id
        WHERE grammar_lessons.id = ?
        """,
        (lesson_id,),
    ).fetchone()


def get_all_grammar_lessons(user=None):
    rows = get_db().execute(
        """
        SELECT grammar_lessons.*, users.fullname AS teacher_name
        FROM grammar_lessons
        JOIN users ON users.id = grammar_lessons.teacher_id
        ORDER BY grammar_lessons.updated_at DESC, grammar_lessons.id DESC
        """
    ).fetchall()
    if not user or user["role"] != "student":
        return rows

    completed_ids = {
        row["lesson_id"]
        for row in get_db().execute(
            "SELECT lesson_id FROM grammar_attempts WHERE user_id = ?",
            (user["id"],),
        ).fetchall()
    }
    lessons = []
    for row in rows:
        lesson = dict(row)
        lesson["is_completed"] = row["id"] in completed_ids
        lessons.append(lesson)
    return lessons


def get_writing_task_or_404(task_id):
    return get_db().execute(
        """
        SELECT writing_tasks.*, users.fullname AS teacher_name
        FROM writing_tasks
        JOIN users ON users.id = writing_tasks.teacher_id
        WHERE writing_tasks.id = ?
        """,
        (task_id,),
    ).fetchone()


def get_all_writing_tasks(user=None):
    rows = get_db().execute(
        """
        SELECT writing_tasks.*, users.fullname AS teacher_name
        FROM writing_tasks
        JOIN users ON users.id = writing_tasks.teacher_id
        ORDER BY writing_tasks.updated_at DESC, writing_tasks.id DESC
        """
    ).fetchall()
    if not user or user["role"] != "student":
        return rows

    submissions = {
        row["task_id"]: row
        for row in get_db().execute(
            "SELECT * FROM writing_submissions WHERE user_id = ?",
            (user["id"],),
        ).fetchall()
    }
    tasks = []
    for row in rows:
        task = dict(row)
        submission = submissions.get(row["id"])
        task["is_completed"] = submission is not None
        task["ai_score"] = submission["ai_score"] if submission else None
        task["teacher_score"] = submission["teacher_score"] if submission else None
        tasks.append(task)
    return tasks


def get_writing_submission(user_id, task_id):
    return get_db().execute(
        "SELECT * FROM writing_submissions WHERE user_id = ? AND task_id = ?",
        (user_id, task_id),
    ).fetchone()


def get_teacher_writing_submissions():
    return get_db().execute(
        """
        SELECT writing_submissions.*, writing_tasks.title AS task_title, users.fullname AS student_name
        FROM writing_submissions
        JOIN writing_tasks ON writing_tasks.id = writing_submissions.task_id
        JOIN users ON users.id = writing_submissions.user_id
        ORDER BY writing_submissions.submitted_at DESC, writing_submissions.id DESC
        """
    ).fetchall()


def get_speaking_lesson_or_404(lesson_id):
    return get_db().execute(
        """
        SELECT speaking_lessons.*, users.fullname AS teacher_name
        FROM speaking_lessons
        JOIN users ON users.id = speaking_lessons.teacher_id
        WHERE speaking_lessons.id = ?
        """,
        (lesson_id,),
    ).fetchone()


def get_all_speaking_lessons(user=None):
    rows = get_db().execute(
        """
        SELECT speaking_lessons.*, users.fullname AS teacher_name
        FROM speaking_lessons
        JOIN users ON users.id = speaking_lessons.teacher_id
        ORDER BY speaking_lessons.updated_at DESC, speaking_lessons.id DESC
        """
    ).fetchall()
    if not user or user["role"] != "student":
        return rows

    attempts = {
        row["lesson_id"]: row
        for row in get_db().execute(
            "SELECT * FROM speaking_attempts WHERE user_id = ?",
            (user["id"],),
        ).fetchall()
    }
    lessons = []
    for row in rows:
        lesson = dict(row)
        attempt = attempts.get(row["id"])
        lesson["is_completed"] = attempt is not None
        lesson["score"] = attempt["score"] if attempt else None
        lessons.append(lesson)
    return lessons


def normalize_practice_words(text):
    return re.findall(r"[a-zA-Z']+", text.lower())


def get_gemini_api_keys():
    keys = [
        os.environ.get("GEMINI_API_KEY1"),
        os.environ.get("GEMINI_API_KEY2"),
        os.environ.get("GEMINI_API_KEY3"),
        os.environ.get("GEMINI_API_KEY"),
    ]
    result = []
    for key in keys:
        if key and key not in result:
            result.append(key)
    return result


def call_gemini_text(prompt):
    api_keys = get_gemini_api_keys()
    if not api_keys:
        return ""

    payload = {
        "contents": [
            {
                "parts": [
                    {"text": prompt}
                ]
            }
        ]
    }
    for api_key in api_keys:
        req = urllib_request.Request(
            f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={api_key}",
            data=json.dumps(payload).encode("utf-8"),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib_request.urlopen(req, timeout=12) as response:
                data = json.loads(response.read().decode("utf-8"))
        except (urllib.error.URLError, TimeoutError, json.JSONDecodeError):
            continue

        candidates = data.get("candidates") or []
        if not candidates:
            continue
        parts = candidates[0].get("content", {}).get("parts", [])
        text = "\n".join(part.get("text", "") for part in parts).strip()
        if text:
            return text
    return ""


def generate_word_notes(passage):
    words = []
    seen = set()
    for word in normalize_practice_words(passage):
        if word not in seen:
            seen.add(word)
            words.append(word)
        if len(words) >= 80:
            break

    prompt = (
        "Return only JSON array. For each English word, give keys word, pronunciation, meaning_vi. "
        f"Words: {', '.join(words)}"
    )
    gemini_text = call_gemini_text(prompt)
    if gemini_text:
        cleaned = gemini_text.strip().removeprefix("```json").removeprefix("```").removesuffix("```").strip()
        try:
            items = json.loads(cleaned)
            if isinstance(items, list):
                return [
                    {
                        "word": str(item.get("word", "")).strip() or word,
                        "pronunciation": str(item.get("pronunciation", "")).strip() or f"/{word}/",
                        "meaning": str(item.get("meaning_vi", item.get("meaning", ""))).strip() or "Dang cap nhat nghia.",
                    }
                    for item, word in zip(items, words)
                ]
        except (json.JSONDecodeError, AttributeError):
            pass

    common_meanings = {
        "hello": "xin chao",
        "name": "ten",
        "school": "truong hoc",
        "student": "hoc sinh",
        "teacher": "giao vien",
        "english": "tieng Anh",
        "learn": "hoc",
        "practice": "luyen tap",
        "today": "hom nay",
        "family": "gia dinh",
        "friend": "ban be",
        "good": "tot",
        "beautiful": "dep",
        "happy": "vui ve",
    }
    return [
        {
            "word": word,
            "pronunciation": f"/{word}/",
            "meaning": common_meanings.get(word, "Chua co nghia tu dong. Cau hinh GEMINI_API_KEY de AI bo sung."),
        }
        for word in words
    ]


def get_word_notes(lesson):
    try:
        notes = json.loads(lesson["word_notes_json"] or "[]")
        return notes if isinstance(notes, list) else []
    except json.JSONDecodeError:
        return []


def evaluate_writing_with_ai(prompt_text, content):
    gemini_prompt = (
        "You are an English writing examiner. Return concise Vietnamese feedback with a score 0-100. "
        f"Writing prompt: {prompt_text}\nStudent answer: {content}"
    )
    gemini_feedback = call_gemini_text(gemini_prompt)
    if gemini_feedback:
        score_match = re.search(r"\b(100|[1-9]?\d)\b", gemini_feedback)
        score = int(score_match.group(1)) if score_match else 75
        return max(0, min(100, score)), gemini_feedback[:1200]

    word_count = len(normalize_practice_words(content))
    score = min(95, max(45, 45 + word_count // 3))
    if word_count < 40:
        feedback = "Bai viet con ngan. Hay viet them y, vi du va cau ket luan de bai ro hon."
    elif word_count < 100:
        feedback = "Bai viet da co y chinh. Nen them tu noi va kiem tra lai ngu phap/cach dung thi."
    else:
        feedback = "Bai viet co do dai tot. Hay tiep tuc chau chuot tu vung, cau phuc va vi du cu the."
    return score, feedback


def clean_json_text(raw_text):
    text = (raw_text or "").strip()
    if text.startswith("```json"):
        text = text[7:].strip()
    elif text.startswith("```"):
        text = text[3:].strip()
    if text.endswith("```"):
        text = text[:-3].strip()
    return text


def get_exam_speaking_topic_or_404(topic_id):
    return get_db().execute(
        """
        SELECT exam_speaking_topics.*, users.fullname AS teacher_name
        FROM exam_speaking_topics
        JOIN users ON users.id = exam_speaking_topics.teacher_id
        WHERE exam_speaking_topics.id = ?
        """,
        (topic_id,),
    ).fetchone()


def get_all_exam_speaking_topics(user=None):
    rows = get_db().execute(
        """
        SELECT exam_speaking_topics.*, users.fullname AS teacher_name
        FROM exam_speaking_topics
        JOIN users ON users.id = exam_speaking_topics.teacher_id
        ORDER BY exam_speaking_topics.updated_at DESC, exam_speaking_topics.id DESC
        """
    ).fetchall()
    if not user or user["role"] != "student":
        return rows

    best_scores = {
        row["topic_id"]: row["best_score"]
        for row in get_db().execute(
            """
            SELECT topic_id, MAX(score) AS best_score
            FROM exam_speaking_sessions
            WHERE user_id = ?
            GROUP BY topic_id
            """,
            (user["id"],),
        ).fetchall()
    }
    topics = []
    for row in rows:
        topic = dict(row)
        topic["best_score"] = best_scores.get(row["id"])
        topic["is_completed"] = topic["best_score"] is not None
        topics.append(topic)
    return topics


def get_exam_speaking_sessions(topic_id=None):
    params = []
    where_sql = ""
    if topic_id:
        where_sql = "WHERE exam_speaking_sessions.topic_id = ?"
        params.append(topic_id)
    return get_db().execute(
        f"""
        SELECT exam_speaking_sessions.*, exam_speaking_topics.title AS topic_title,
               users.fullname AS student_name
        FROM exam_speaking_sessions
        JOIN exam_speaking_topics ON exam_speaking_topics.id = exam_speaking_sessions.topic_id
        JOIN users ON users.id = exam_speaking_sessions.user_id
        {where_sql}
        ORDER BY exam_speaking_sessions.completed_at DESC, exam_speaking_sessions.id DESC
        """,
        params,
    ).fetchall()


def build_exam_speaking_ai_reply(topic, conversation, turn_index):
    fallback_questions = [
        "Hello, nice to meet you. What is your name and where are you from?",
        f"Good. Now let's talk about {topic['title']}. Can you share your first idea?",
        "Why do you think that is important?",
        "Can you give me one example from your real life?",
        "Thank you. Please give a short final answer about this topic.",
    ]

    prompt = f"""
You are a friendly English speaking examiner for Vietnamese students.
Test topic: {topic['title']}
Teacher instruction/topic detail: {topic['topic_prompt']}
Level: {topic['level']}
Opening question from teacher, if any: {topic['opening_question'] or ''}
Conversation so far JSON:
{json.dumps(conversation, ensure_ascii=False)}

Return only JSON with keys:
ai_text: one short spoken English response/question, 1-2 sentences.
finished: boolean true only after enough topic discussion, usually after 5 student answers.

Rules:
- First turn: greet, ask name/location/warm-up.
- After warm-up, ask directly about the teacher topic.
- Do not give score until final endpoint.
- Keep language natural and not too long.
Current turn index: {turn_index}
"""
    text = call_gemini_text(prompt)
    if text:
        try:
            data = json.loads(clean_json_text(text))
            ai_text = str(data.get("ai_text", "")).strip()
            if ai_text:
                return {
                    "ai_text": ai_text,
                    "finished": bool(data.get("finished", False)),
                    "source": "gemini",
                }
        except (json.JSONDecodeError, AttributeError):
            pass

    index = min(turn_index, len(fallback_questions) - 1)
    return {
        "ai_text": fallback_questions[index],
        "finished": turn_index >= len(fallback_questions) - 1,
        "source": "fallback",
    }


def evaluate_exam_speaking(topic, conversation):
    prompt = f"""
You are an English speaking examiner. Evaluate a Vietnamese student's spoken English conversation.
Topic: {topic['title']}
Teacher topic detail: {topic['topic_prompt']}
Level: {topic['level']}
Conversation JSON:
{json.dumps(conversation, ensure_ascii=False)}

Return only JSON:
score: integer 0-100
feedback_vi: concise Vietnamese feedback
pronunciation_feedback_vi: Vietnamese feedback about pronunciation/clarity
unclear_words: array of objects with word, issue_vi, suggestion_vi

Focus on topic relevance, fluency, grammar, vocabulary, pronunciation clarity, unclear words and wrong words.
"""
    text = call_gemini_text(prompt)
    if text:
        try:
            data = json.loads(clean_json_text(text))
            score = int(data.get("score", 0))
            unclear_words = data.get("unclear_words") or []
            if not isinstance(unclear_words, list):
                unclear_words = []
            return {
                "score": max(0, min(100, score)),
                "feedback": str(data.get("feedback_vi", "")).strip()[:1600],
                "pronunciation_feedback": str(data.get("pronunciation_feedback_vi", "")).strip()[:1600],
                "unclear_words": unclear_words[:20],
                "source": "gemini",
            }
        except (json.JSONDecodeError, TypeError, ValueError, AttributeError):
            pass

    student_text = " ".join(item.get("student", "") for item in conversation if isinstance(item, dict))
    word_count = len(normalize_practice_words(student_text))
    score = min(88, max(45, 45 + word_count // 4))
    return {
        "score": score,
        "feedback": "AI fallback: Cau tra loi da duoc ghi nhan. Hay tra loi dai hon, bam sat chu de hon va them vi du cu the.",
        "pronunciation_feedback": "Chua co Gemini key nen he thong chi luu transcript. Hay cau hinh GEMINI_API_KEY1/2/3 de phan tich phat am chi tiet.",
        "unclear_words": [],
        "source": "fallback",
    }


def exam_docx_is_allowed(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() == "docx"


def parse_exam_datetime(value):
    try:
        return datetime.strptime(value, "%Y-%m-%dT%H:%M")
    except (TypeError, ValueError):
        return None


def get_exam_room_status(room):
    start_at = parse_exam_datetime(room["start_at"])
    if not start_at:
        return "unknown"
    now = datetime.now()
    duration = int(room["duration_minutes"] or 45)
    end_at = start_at.timestamp() + duration * 60
    if now.timestamp() < start_at.timestamp():
        return "waiting"
    if now.timestamp() > end_at:
        return "closed"
    return "open"


def normalize_option_text(text):
    return re.sub(r"^[A-Da-d][\.\)]\s*", "", text or "").strip()


def paragraph_has_underline(paragraph):
    return any((run.font.underline or False) and run.text.strip() for run in paragraph.runs)


def read_docx_paragraphs(file_path):
    if Document is None:
        raise RuntimeError("python-docx is not installed")
    document = Document(file_path)
    paragraphs = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "text": text,
                    "underlined": paragraph_has_underline(paragraph),
                }
            )
    return paragraphs


def finalize_exam_question(parsed, question_type, passage, question_text, options, correct_answer):
    if not question_text or not all(options.get(key) for key in ("A", "B", "C", "D")):
        return
    parsed.append(
        {
            "question_type": question_type or "LN1",
            "passage": passage.strip(),
            "question_text": question_text.strip(),
            "option_a": options.get("A", "").strip(),
            "option_b": options.get("B", "").strip(),
            "option_c": options.get("C", "").strip(),
            "option_d": options.get("D", "").strip(),
            "correct_answer": correct_answer if correct_answer in {"A", "B", "C", "D"} else "A",
        }
    )


def parse_exam_docx(file_path):
    paragraphs = read_docx_paragraphs(file_path)
    parsed = []
    current_type = None
    passage_parts = []
    question_text = ""
    options = {}
    correct_answer = ""
    ln2_has_question = False

    def reset_question():
        return "", {}, ""

    for item in paragraphs:
        text = item["text"]
        marker_match = re.match(r"^\((LN[12])\)\s*(.*)$", text, re.IGNORECASE)
        if marker_match:
            finalize_exam_question(parsed, current_type, "\n".join(passage_parts), question_text, options, correct_answer)
            current_type = marker_match.group(1).upper()
            rest = marker_match.group(2).strip()
            passage_parts = []
            question_text, options, correct_answer = reset_question()
            ln2_has_question = False
            if rest:
                if current_type == "LN2":
                    passage_parts.append(rest)
                else:
                    question_text = rest
            continue

        if current_type is None:
            continue

        option_match = re.match(r"^([A-Da-d])[\.\)]\s*(.+)$", text)
        if option_match:
            letter = option_match.group(1).upper()
            options[letter] = normalize_option_text(text)
            if item["underlined"]:
                correct_answer = letter
            if all(options.get(key) for key in ("A", "B", "C", "D")):
                finalize_exam_question(parsed, current_type, "\n".join(passage_parts), question_text, options, correct_answer)
                question_text, options, correct_answer = reset_question()
                if current_type == "LN1":
                    passage_parts = []
                ln2_has_question = False
            continue

        if current_type == "LN2" and not question_text:
            looks_like_question = bool(re.match(r"^(\d+[\.\)]|Câu|Question)\s+", text, re.IGNORECASE)) or "?" in text
            if not ln2_has_question and not looks_like_question:
                passage_parts.append(text)
                continue
            ln2_has_question = True

        if question_text and not options:
            question_text = f"{question_text}\n{text}".strip()
        else:
            question_text = text

    finalize_exam_question(parsed, current_type, "\n".join(passage_parts), question_text, options, correct_answer)
    return parsed


def get_exam_room_or_404(room_id):
    return get_db().execute(
        """
        SELECT exam_rooms.*, users.fullname AS teacher_name
        FROM exam_rooms
        JOIN users ON users.id = exam_rooms.teacher_id
        WHERE exam_rooms.id = ?
        """,
        (room_id,),
    ).fetchone()


def get_exam_room_questions(room_id):
    return get_db().execute(
        """
        SELECT *
        FROM exam_room_questions
        WHERE room_id = ?
        ORDER BY position ASC, id ASC
        """,
        (room_id,),
    ).fetchall()


def get_all_exam_rooms(user=None):
    rooms = get_db().execute(
        """
        SELECT exam_rooms.*, users.fullname AS teacher_name,
               COUNT(DISTINCT exam_room_registrations.id) AS registration_count,
               COUNT(DISTINCT exam_room_submissions.id) AS submission_count
        FROM exam_rooms
        JOIN users ON users.id = exam_rooms.teacher_id
        LEFT JOIN exam_room_registrations ON exam_room_registrations.room_id = exam_rooms.id
        LEFT JOIN exam_room_submissions ON exam_room_submissions.room_id = exam_rooms.id
        GROUP BY exam_rooms.id
        ORDER BY exam_rooms.start_at DESC, exam_rooms.id DESC
        """
    ).fetchall()
    if not user or user["role"] != "student":
        return rooms

    registrations = {
        row["room_id"]
        for row in get_db().execute(
            "SELECT room_id FROM exam_room_registrations WHERE user_id = ?",
            (user["id"],),
        ).fetchall()
    }
    submissions = {
        row["room_id"]: row
        for row in get_db().execute(
            "SELECT room_id, score FROM exam_room_submissions WHERE user_id = ?",
            (user["id"],),
        ).fetchall()
    }
    result = []
    for row in rooms:
        item = dict(row)
        item["is_registered"] = row["id"] in registrations
        item["is_submitted"] = row["id"] in submissions
        item["score"] = submissions[row["id"]]["score"] if row["id"] in submissions else None
        item["status"] = get_exam_room_status(row)
        result.append(item)
    return result


def get_preview_questions_from_form():
    questions = []
    try:
        total = int(request.form.get("question_count", "0"))
    except ValueError:
        total = 0
    for index in range(1, total + 1):
        question_type = request.form.get(f"q_{index}_type", "LN1").strip().upper()
        passage = request.form.get(f"q_{index}_passage", "").strip()
        question_text = request.form.get(f"q_{index}_text", "").strip()
        option_a = request.form.get(f"q_{index}_a", "").strip()
        option_b = request.form.get(f"q_{index}_b", "").strip()
        option_c = request.form.get(f"q_{index}_c", "").strip()
        option_d = request.form.get(f"q_{index}_d", "").strip()
        correct_answer = request.form.get(f"q_{index}_correct", "").strip().upper()
        if not all([question_text, option_a, option_b, option_c, option_d]) or correct_answer not in {"A", "B", "C", "D"}:
            continue
        questions.append(
            {
                "question_type": question_type if question_type in {"LN1", "LN2"} else "LN1",
                "passage": passage,
                "question_text": question_text,
                "option_a": option_a,
                "option_b": option_b,
                "option_c": option_c,
                "option_d": option_d,
                "correct_answer": correct_answer,
            }
        )
    return questions


def get_exam_room_submission(user_id, room_id):
    return get_db().execute(
        "SELECT * FROM exam_room_submissions WHERE user_id = ? AND room_id = ?",
        (user_id, room_id),
    ).fetchone()


def student_is_registered_for_room(user_id, room_id):
    return get_db().execute(
        "SELECT id FROM exam_room_registrations WHERE user_id = ? AND room_id = ?",
        (user_id, room_id),
    ).fetchone() is not None


def get_student_home_data(user):
    db = get_db()
    current_score_row = db.execute(
        "SELECT COALESCE(MAX(score), 0) AS best_score, COUNT(*) AS exam_count FROM exam_results WHERE user_id = ?",
        (user["id"],),
    ).fetchone()
    best_score = current_score_row["best_score"] or 918
    exam_count = current_score_row["exam_count"] or 3
    completed_activity_count = current_score_row["exam_count"] or 3
    progress = min(100, completed_activity_count * 10)

    return {
        "progress": progress,
        "best_score": best_score,
        "exam_count": exam_count,
        "study_minutes": 420 + completed_activity_count * 15,
        "lesson_done": completed_activity_count,
        "skills": [
            {"name": "Nghe", "value": 82},
            {"name": "Nói", "value": 74},
            {"name": "Đọc", "value": 88},
            {"name": "Viết", "value": 69},
            {"name": "Từ vựng", "value": 91},
        ],
    }

@app.teardown_appcontext
def close_connection(exception):
    db = getattr(g, "_database", None)
    if db is not None:
        db.close()


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/healthz")
def healthz():
    return jsonify({"status": "ok"})


@app.route("/img/<path:filename>")
def image_asset(filename):
    return send_from_directory(BASE_DIR / "img", filename)


@app.route("/register", methods=["GET", "POST"])
def register():
    init_db()
    if request.method == "POST":
        fullname = request.form.get("fullname", "").strip()
        username = request.form.get("username", "").strip() or None
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        class_name = request.form.get("class", "").strip() or None

        if not fullname or not email or not password:
            flash("Vui lòng điền đầy đủ họ tên, email và mật khẩu.")
            return redirect(url_for("register"))

        db = get_db()
        user_exists = db.execute(
            "SELECT id FROM users WHERE email = ? OR username = ?",
            (email, username),
        ).fetchone()
        if user_exists:
            flash("Email hoặc tên đăng nhập đã tồn tại.")
            return redirect(url_for("register"))

        db.execute(
            "INSERT INTO users (fullname, username, email, password_hash, role, class) VALUES (?, ?, ?, ?, ?, ?)",
            (fullname, username, email, generate_password_hash(password), "student", class_name),
        )
        db.commit()
        flash("Đăng ký thành công. Bạn có thể đăng nhập ngay bây giờ.")
        return redirect(url_for("login"))

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    init_db()
    if request.method == "POST":
        identifier = request.form.get("identifier", "").strip()
        password = request.form.get("password", "")

        if not identifier or not password:
            flash("Vui lòng nhập email/tên đăng nhập và mật khẩu.")
            return redirect(url_for("login"))

        db = get_db()
        user = db.execute(
            "SELECT * FROM users WHERE email = ? OR username = ?",
            (identifier.lower(), identifier),
        ).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"] = user["id"]
            session["role"] = user["role"]
            session["username"] = user["username"] or user["email"]
            flash("Đăng nhập thành công.")
            return redirect(url_for("home"))

        flash("Email/tên đăng nhập hoặc mật khẩu không đúng.")
        return redirect(url_for("login"))

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("Bạn đã đăng xuất.")
    return redirect(url_for("index"))


@app.route("/home")
def home():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if session.get("role") == "teacher":
        return render_template("dashboard_teacher.html", user=user)
    return render_template(
        "dashboard_student.html",
        user=user,
        home_data=get_student_home_data(user),
        game_packages=get_all_quiz_packages(user),
    )


@app.route("/skills")
def skills():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    template_name = "skills_teacher.html" if user["role"] == "teacher" else "skills_student.html"
    return render_template(template_name, user=user)


@app.route("/exams/speaking-ai", methods=["GET", "POST"])
def exam_speaking_ai():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if request.method == "POST":
        if user["role"] != "teacher":
            flash("Chi giao vien moi duoc tao chu de Speaking AI Test.")
            return redirect(url_for("exam_speaking_ai"))

        title = request.form.get("title", "").strip()
        topic_prompt = request.form.get("topic_prompt", "").strip()
        level = request.form.get("level", "A2-B1").strip() or "A2-B1"
        opening_question = request.form.get("opening_question", "").strip() or None
        if not title or not topic_prompt:
            flash("Vui long nhap tieu de va noi dung chu de.")
            return redirect(url_for("exam_speaking_ai"))

        get_db().execute(
            """
            INSERT INTO exam_speaking_topics
                (teacher_id, title, topic_prompt, level, opening_question)
            VALUES (?, ?, ?, ?, ?)
            """,
            (user["id"], title, topic_prompt, level, opening_question),
        )
        get_db().commit()
        flash("Da tao chu de Speaking AI Test.")
        return redirect(url_for("exam_speaking_ai"))

    template_name = "exam_speaking_teacher.html" if user["role"] == "teacher" else "exam_speaking_student.html"
    return render_template(
        template_name,
        user=user,
        topics=get_all_exam_speaking_topics(user),
        sessions=get_exam_speaking_sessions() if user["role"] == "teacher" else [],
        edit_topic=None,
    )


@app.route("/exams/speaking-ai/<int:topic_id>")
def exam_speaking_detail(topic_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    topic = get_exam_speaking_topic_or_404(topic_id)
    if not topic:
        flash("Khong tim thay chu de Speaking AI Test.")
        return redirect(url_for("exam_speaking_ai"))

    if user["role"] != "student":
        return redirect(url_for("exam_speaking_ai"))

    recent_sessions = get_db().execute(
        """
        SELECT *
        FROM exam_speaking_sessions
        WHERE user_id = ? AND topic_id = ?
        ORDER BY completed_at DESC, id DESC
        LIMIT 5
        """,
        (user["id"], topic_id),
    ).fetchall()
    return render_template(
        "exam_speaking_detail.html",
        user=user,
        topic=topic,
        recent_sessions=recent_sessions,
    )


@app.route("/exams/speaking-ai/<int:topic_id>/edit", methods=["GET", "POST"])
def exam_speaking_edit(topic_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc chinh sua chu de Speaking AI Test.")
        return redirect(url_for("exam_speaking_ai"))

    topic = get_exam_speaking_topic_or_404(topic_id)
    if not topic:
        flash("Khong tim thay chu de Speaking AI Test.")
        return redirect(url_for("exam_speaking_ai"))

    if request.method == "POST":
        title = request.form.get("title", "").strip()
        topic_prompt = request.form.get("topic_prompt", "").strip()
        level = request.form.get("level", "A2-B1").strip() or "A2-B1"
        opening_question = request.form.get("opening_question", "").strip() or None
        if not title or not topic_prompt:
            flash("Vui long nhap tieu de va noi dung chu de.")
            return redirect(url_for("exam_speaking_edit", topic_id=topic_id))

        get_db().execute(
            """
            UPDATE exam_speaking_topics
            SET title = ?, topic_prompt = ?, level = ?, opening_question = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (title, topic_prompt, level, opening_question, topic_id),
        )
        get_db().commit()
        flash("Da cap nhat chu de Speaking AI Test.")
        return redirect(url_for("exam_speaking_ai"))

    return render_template(
        "exam_speaking_teacher.html",
        user=user,
        topics=get_all_exam_speaking_topics(user),
        sessions=get_exam_speaking_sessions(),
        edit_topic=topic,
    )


@app.route("/exams/speaking-ai/<int:topic_id>/delete", methods=["POST"])
def exam_speaking_delete(topic_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc xoa chu de Speaking AI Test.")
        return redirect(url_for("exam_speaking_ai"))

    get_db().execute("DELETE FROM exam_speaking_sessions WHERE topic_id = ?", (topic_id,))
    get_db().execute("DELETE FROM exam_speaking_topics WHERE id = ?", (topic_id,))
    get_db().commit()
    flash("Da xoa chu de Speaking AI Test.")
    return redirect(url_for("exam_speaking_ai"))


@app.route("/api/exams/speaking-ai/<int:topic_id>/turn", methods=["POST"])
def exam_speaking_turn(topic_id):
    user = get_current_user()
    if not user or user["role"] != "student":
        return jsonify({"error": "unauthorized"}), 401

    topic = get_exam_speaking_topic_or_404(topic_id)
    if not topic:
        return jsonify({"error": "topic_not_found"}), 404

    data = request.get_json(silent=True) or {}
    conversation = data.get("conversation") or []
    if not isinstance(conversation, list):
        conversation = []
    try:
        turn_index = int(data.get("turn_index", len(conversation)))
    except (TypeError, ValueError):
        turn_index = len(conversation)

    reply = build_exam_speaking_ai_reply(topic, conversation, turn_index)
    return jsonify(reply)


@app.route("/api/exams/speaking-ai/<int:topic_id>/finish", methods=["POST"])
def exam_speaking_finish(topic_id):
    user = get_current_user()
    if not user or user["role"] != "student":
        return jsonify({"error": "unauthorized"}), 401

    topic = get_exam_speaking_topic_or_404(topic_id)
    if not topic:
        return jsonify({"error": "topic_not_found"}), 404

    data = request.get_json(silent=True) or {}
    conversation = data.get("conversation") or []
    if not isinstance(conversation, list):
        conversation = []
    result = evaluate_exam_speaking(topic, conversation)

    get_db().execute(
        """
        INSERT INTO exam_speaking_sessions
            (topic_id, user_id, conversation_json, score, feedback,
             pronunciation_feedback, unclear_words_json, status, completed_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, 'completed', CURRENT_TIMESTAMP)
        """,
        (
            topic_id,
            user["id"],
            json.dumps(conversation, ensure_ascii=False),
            result["score"],
            result["feedback"],
            result["pronunciation_feedback"],
            json.dumps(result["unclear_words"], ensure_ascii=False),
        ),
    )
    get_db().commit()
    return jsonify(result)


@app.route("/exams/rooms", methods=["GET", "POST"])
def exam_rooms():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if request.method == "POST":
        if user["role"] != "teacher":
            flash("Chi giao vien moi duoc tao phong thi.")
            return redirect(url_for("exam_rooms"))
        if Document is None:
            flash("Chua cai python-docx. Hay chay: pip install -r requirements.txt")
            return redirect(url_for("exam_rooms"))

        title = request.form.get("title", "").strip()
        description = request.form.get("description", "").strip()
        start_at = request.form.get("start_at", "").strip()
        try:
            duration_minutes = int(request.form.get("duration_minutes", "45"))
        except ValueError:
            duration_minutes = 45
        duration_minutes = max(5, min(240, duration_minutes))
        exam_file = request.files.get("exam_file")

        if not title or not start_at or not exam_file or not exam_file.filename:
            flash("Vui long nhap tieu de, gio mo phong va upload file Word .docx.")
            return redirect(url_for("exam_rooms"))
        if not parse_exam_datetime(start_at):
            flash("Gio mo phong khong hop le.")
            return redirect(url_for("exam_rooms"))
        if not exam_docx_is_allowed(exam_file.filename):
            flash("Phong thi hien ho tro file .docx de doc dap an gach chan. Hay luu Word sang .docx.")
            return redirect(url_for("exam_rooms"))

        original_name = secure_filename(exam_file.filename)
        saved_name = f"room_{user['id']}_{int(time.time())}_{original_name}"
        file_path = EXAM_FILE_DIR / saved_name
        exam_file.save(file_path)
        try:
            questions = parse_exam_docx(file_path)
        except Exception:
            flash("Khong doc duoc file Word. Hay kiem tra dinh dang (LN1)/(LN2), dap an ABCD va dap an dung duoc gach chan.")
            return redirect(url_for("exam_rooms"))
        if not questions:
            flash("Chua tach duoc cau hoi nao. Hay kiem tra file Word co marker (LN1)/(LN2) va dap an ABCD.")
            return redirect(url_for("exam_rooms"))

        return render_template(
            "exam_room_preview.html",
            user=user,
            title=title,
            description=description,
            start_at=start_at,
            duration_minutes=duration_minutes,
            source_file_name=exam_file.filename,
            questions=questions,
        )

    template_name = "exam_rooms_teacher.html" if user["role"] == "teacher" else "exam_rooms_student.html"
    return render_template(template_name, user=user, rooms=get_all_exam_rooms(user))


@app.route("/exams/rooms/save", methods=["POST"])
def exam_room_save():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc luu phong thi.")
        return redirect(url_for("exam_rooms"))

    title = request.form.get("title", "").strip()
    description = request.form.get("description", "").strip()
    start_at = request.form.get("start_at", "").strip()
    source_file_name = request.form.get("source_file_name", "").strip()
    try:
        duration_minutes = int(request.form.get("duration_minutes", "45"))
    except ValueError:
        duration_minutes = 45
    duration_minutes = max(5, min(240, duration_minutes))
    questions = get_preview_questions_from_form()

    if not title or not parse_exam_datetime(start_at) or not questions:
        flash("Phong thi chua hop le. Can co tieu de, gio mo phong va it nhat 1 cau hoi day du.")
        return redirect(url_for("exam_rooms"))

    db = get_db()
    cursor = db.execute(
        """
        INSERT INTO exam_rooms
            (teacher_id, title, description, start_at, duration_minutes, source_file_name, total_questions)
        VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (user["id"], title, description, start_at, duration_minutes, source_file_name, len(questions)),
    )
    room_id = cursor.lastrowid
    for position, question in enumerate(questions, start=1):
        db.execute(
            """
            INSERT INTO exam_room_questions
                (room_id, position, question_type, passage, question_text,
                 option_a, option_b, option_c, option_d, correct_answer)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                room_id, position, question["question_type"], question["passage"],
                question["question_text"], question["option_a"], question["option_b"],
                question["option_c"], question["option_d"], question["correct_answer"],
            ),
        )
    db.commit()
    flash("Da luu phong thi.")
    return redirect(url_for("exam_rooms"))


@app.route("/exams/rooms/<int:room_id>")
def exam_room_detail(room_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    room = get_exam_room_or_404(room_id)
    if not room:
        flash("Khong tim thay phong thi.")
        return redirect(url_for("exam_rooms"))
    questions = get_exam_room_questions(room_id)

    if user["role"] == "teacher":
        submissions = get_db().execute(
            """
            SELECT exam_room_submissions.*, users.fullname AS student_name
            FROM exam_room_submissions
            JOIN users ON users.id = exam_room_submissions.user_id
            WHERE exam_room_submissions.room_id = ?
            ORDER BY exam_room_submissions.score DESC, exam_room_submissions.submitted_at ASC
            """,
            (room_id,),
        ).fetchall()
        return render_template(
            "exam_room_teacher_detail.html",
            user=user,
            room=room,
            questions=questions,
            submissions=submissions,
        )

    is_registered = student_is_registered_for_room(user["id"], room_id)
    submission = get_exam_room_submission(user["id"], room_id)
    status = get_exam_room_status(room)
    if not is_registered:
        flash("Ban can dang ky phong thi truoc khi vao lam.")
        return redirect(url_for("exam_rooms"))
    return render_template(
        "exam_room_take.html",
        user=user,
        room=room,
        questions=questions,
        status=status,
        submission=submission,
    )


@app.route("/exams/rooms/<int:room_id>/register", methods=["POST"])
def exam_room_register(room_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "student":
        flash("Chi hoc sinh moi dang ky phong thi.")
        return redirect(url_for("exam_rooms"))

    room = get_exam_room_or_404(room_id)
    if not room:
        flash("Khong tim thay phong thi.")
        return redirect(url_for("exam_rooms"))
    get_db().execute(
        """
        INSERT OR IGNORE INTO exam_room_registrations (room_id, user_id)
        VALUES (?, ?)
        """,
        (room_id, user["id"]),
    )
    get_db().commit()
    flash("Da dang ky phong thi. Dung gio mo phong thi moi vao lam duoc.")
    return redirect(url_for("exam_rooms"))


@app.route("/exams/rooms/<int:room_id>/submit", methods=["POST"])
def exam_room_submit(room_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "student":
        flash("Chi hoc sinh moi nop bai thi.")
        return redirect(url_for("exam_rooms"))

    room = get_exam_room_or_404(room_id)
    if not room:
        flash("Khong tim thay phong thi.")
        return redirect(url_for("exam_rooms"))
    if not student_is_registered_for_room(user["id"], room_id):
        flash("Ban can dang ky phong thi truoc khi nop bai.")
        return redirect(url_for("exam_rooms"))
    if get_exam_room_status(room) != "open":
        flash("Phong thi chua mo hoac da het gio.")
        return redirect(url_for("exam_room_detail", room_id=room_id))
    if get_exam_room_submission(user["id"], room_id):
        flash("Ban da nop bai phong thi nay.")
        return redirect(url_for("exam_room_detail", room_id=room_id))

    questions = get_exam_room_questions(room_id)
    answers = {}
    correct_count = 0
    for question in questions:
        selected = request.form.get(f"answer_{question['id']}", "").strip().upper()
        if selected in {"A", "B", "C", "D"}:
            answers[str(question["id"])] = selected
            if selected == question["correct_answer"]:
                correct_count += 1
    total_questions = len(questions)
    score = round((correct_count / total_questions) * 10, 2) if total_questions else 0
    get_db().execute(
        """
        INSERT INTO exam_room_submissions
            (room_id, user_id, answers_json, correct_count, total_questions, score, submitted_at)
        VALUES (?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        """,
        (room_id, user["id"], json.dumps(answers), correct_count, total_questions, score),
    )
    get_db().execute(
        """
        INSERT INTO exam_results (user_id, room_name, score, taken_at)
        VALUES (?, ?, ?, CURRENT_TIMESTAMP)
        """,
        (user["id"], room["title"], int(round(score * 100))),
    )
    get_db().commit()
    flash("Da nop bai thi.")
    return redirect(url_for("exam_room_detail", room_id=room_id))


@app.route("/exams/rooms/<int:room_id>/delete", methods=["POST"])
def exam_room_delete(room_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc xoa phong thi.")
        return redirect(url_for("exam_rooms"))
    db = get_db()
    db.execute("DELETE FROM exam_room_submissions WHERE room_id = ?", (room_id,))
    db.execute("DELETE FROM exam_room_registrations WHERE room_id = ?", (room_id,))
    db.execute("DELETE FROM exam_room_questions WHERE room_id = ?", (room_id,))
    db.execute("DELETE FROM exam_rooms WHERE id = ?", (room_id,))
    db.commit()
    flash("Da xoa phong thi.")
    return redirect(url_for("exam_rooms"))


@app.route("/skills/listening", methods=["GET", "POST"])
def listening_lessons():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if request.method == "POST":
        if user["role"] != "teacher":
            flash("Chỉ giáo viên mới được tạo bài nghe.")
            return redirect(url_for("listening_lessons"))

        title = request.form.get("title", "").strip()
        passage = request.form.get("passage", "").strip()
        quiz_question = request.form.get("quiz_question", "").strip()
        option_a = request.form.get("option_a", "").strip()
        option_b = request.form.get("option_b", "").strip()
        option_c = request.form.get("option_c", "").strip()
        option_d = request.form.get("option_d", "").strip()
        correct_answer = request.form.get("correct_answer", "").strip().upper()

        if not all([title, passage, quiz_question, option_a, option_b, option_c, option_d]):
            flash("Vui lòng nhập đủ tiêu đề, đoạn nghe và câu hỏi ABCD.")
            return redirect(url_for("listening_lessons"))
        if correct_answer not in {"A", "B", "C", "D"}:
            flash("Vui lòng tick đáp án đúng.")
            return redirect(url_for("listening_lessons"))

        get_db().execute(
            """
            INSERT INTO listening_lessons
                (teacher_id, title, passage, quiz_question, option_a, option_b, option_c, option_d, correct_answer)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (user["id"], title, passage, quiz_question, option_a, option_b, option_c, option_d, correct_answer),
        )
        get_db().commit()
        flash("Đã tạo bài nghe.")
        return redirect(url_for("listening_lessons"))

    template_name = "listening_teacher.html" if user["role"] == "teacher" else "listening_student.html"
    return render_template(
        template_name,
        user=user,
        lessons=get_all_listening_lessons(user),
        edit_lesson=None,
    )


@app.route("/skills/listening/<int:lesson_id>", methods=["GET", "POST"])
def listening_detail(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    lesson = get_listening_lesson_or_404(lesson_id)
    if not lesson:
        flash("Không tìm thấy bài nghe.")
        return redirect(url_for("listening_lessons"))

    attempt = get_learning_attempt("listening_attempts", user["id"], lesson_id) if user["role"] == "student" else None
    if request.method == "POST":
        if user["role"] != "student":
            flash("Chỉ học sinh mới nộp bài nghe.")
            return redirect(url_for("listening_detail", lesson_id=lesson_id))

        selected_answer = request.form.get("answer", "").strip().upper()
        if selected_answer not in {"A", "B", "C", "D"}:
            flash("Vui lòng chọn một đáp án.")
            return redirect(url_for("listening_detail", lesson_id=lesson_id))

        score = 100 if selected_answer == lesson["correct_answer"] else 0
        get_db().execute(
            """
            INSERT INTO listening_attempts (user_id, lesson_id, selected_answer, score, completed_at)
            VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id, lesson_id)
            DO UPDATE SET selected_answer = excluded.selected_answer,
                          score = excluded.score,
                          completed_at = CURRENT_TIMESTAMP
            """,
            (user["id"], lesson_id, selected_answer, score),
        )
        get_db().commit()
        flash("Đã nộp bài nghe.")
        attempt = get_learning_attempt("listening_attempts", user["id"], lesson_id)

    return render_template("listening_detail.html", user=user, lesson=lesson, attempt=attempt)


@app.route("/skills/listening/<int:lesson_id>/edit", methods=["GET", "POST"])
def listening_edit(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chỉ giáo viên mới được chỉnh sửa bài nghe.")
        return redirect(url_for("listening_lessons"))

    lesson = get_listening_lesson_or_404(lesson_id)
    if not lesson:
        flash("Không tìm thấy bài nghe.")
        return redirect(url_for("listening_lessons"))

    if request.method == "POST":
        title = request.form.get("title", "").strip()
        passage = request.form.get("passage", "").strip()
        quiz_question = request.form.get("quiz_question", "").strip()
        option_a = request.form.get("option_a", "").strip()
        option_b = request.form.get("option_b", "").strip()
        option_c = request.form.get("option_c", "").strip()
        option_d = request.form.get("option_d", "").strip()
        correct_answer = request.form.get("correct_answer", "").strip().upper()

        if not all([title, passage, quiz_question, option_a, option_b, option_c, option_d]):
            flash("Vui lòng nhập đủ tiêu đề, đoạn nghe và câu hỏi ABCD.")
            return redirect(url_for("listening_edit", lesson_id=lesson_id))
        if correct_answer not in {"A", "B", "C", "D"}:
            flash("Vui lòng tick đáp án đúng.")
            return redirect(url_for("listening_edit", lesson_id=lesson_id))

        get_db().execute(
            """
            UPDATE listening_lessons
            SET title = ?, passage = ?, quiz_question = ?, option_a = ?, option_b = ?,
                option_c = ?, option_d = ?, correct_answer = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (title, passage, quiz_question, option_a, option_b, option_c, option_d, correct_answer, lesson_id),
        )
        get_db().commit()
        flash("Đã cập nhật bài nghe.")
        return redirect(url_for("listening_lessons"))

    return render_template(
        "listening_teacher.html",
        user=user,
        lessons=get_all_listening_lessons(user),
        edit_lesson=lesson,
    )


@app.route("/skills/listening/<int:lesson_id>/delete", methods=["POST"])
def listening_delete(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chỉ giáo viên mới được xóa bài nghe.")
        return redirect(url_for("listening_lessons"))

    get_db().execute("DELETE FROM listening_attempts WHERE lesson_id = ?", (lesson_id,))
    get_db().execute("DELETE FROM listening_lessons WHERE id = ?", (lesson_id,))
    get_db().commit()
    flash("Đã xóa bài nghe.")
    return redirect(url_for("listening_lessons"))


@app.route("/skills/grammar", methods=["GET", "POST"])
def grammar_lessons():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if request.method == "POST":
        if user["role"] != "teacher":
            flash("Chỉ giáo viên mới được tạo bài ngữ pháp.")
            return redirect(url_for("grammar_lessons"))

        title = request.form.get("title", "").strip()
        content = request.form.get("content", "").strip()
        material_type = request.form.get("material_type", "file").strip()
        material_url = request.form.get("material_url", "").strip()
        material_file = request.files.get("material_file")
        quiz_question = request.form.get("quiz_question", "").strip()
        option_a = request.form.get("option_a", "").strip()
        option_b = request.form.get("option_b", "").strip()
        option_c = request.form.get("option_c", "").strip()
        option_d = request.form.get("option_d", "").strip()
        correct_answer = request.form.get("correct_answer", "").strip().upper()

        if not all([title, content, quiz_question, option_a, option_b, option_c, option_d]):
            flash("Vui lòng nhập đủ tiêu đề, nội dung và câu hỏi ABCD.")
            return redirect(url_for("grammar_lessons"))
        if correct_answer not in {"A", "B", "C", "D"}:
            flash("Vui lòng tick đáp án đúng.")
            return redirect(url_for("grammar_lessons"))

        material_path = None
        material_name = None
        if material_type == "link":
            if not material_url:
                flash("Vui lòng nhập link tài liệu.")
                return redirect(url_for("grammar_lessons"))
            material_url = normalize_embed_url(material_url)
        elif material_type == "file":
            if not material_file or not material_file.filename:
                flash("Vui lòng upload file tài liệu.")
                return redirect(url_for("grammar_lessons"))
            if not material_file_is_allowed(material_file.filename):
                flash("File tài liệu chưa đúng định dạng cho phép.")
                return redirect(url_for("grammar_lessons"))
            original_name = secure_filename(material_file.filename)
            saved_name = f"grammar_file_{user['id']}_{int(time.time())}_{original_name}"
            material_file.save(SKILL_FILE_DIR / saved_name)
            material_path = f"uploads/skills/files/{saved_name}"
            material_name = material_file.filename
            material_url = ""
        else:
            flash("Loại tài liệu không hợp lệ.")
            return redirect(url_for("grammar_lessons"))

        get_db().execute(
            """
            INSERT INTO grammar_lessons
                (teacher_id, title, content, material_type, material_url, material_path,
                 material_name, quiz_question, option_a, option_b, option_c, option_d, correct_answer)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                user["id"], title, content, material_type, material_url or None, material_path,
                material_name, quiz_question, option_a, option_b, option_c, option_d, correct_answer,
            ),
        )
        get_db().commit()
        flash("Đã tạo bài ngữ pháp.")
        return redirect(url_for("grammar_lessons"))

    template_name = "grammar_teacher.html" if user["role"] == "teacher" else "grammar_student.html"
    return render_template(
        template_name,
        user=user,
        lessons=get_all_grammar_lessons(user),
        edit_lesson=None,
    )


@app.route("/skills/grammar/<int:lesson_id>", methods=["GET", "POST"])
def grammar_detail(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    lesson = get_grammar_lesson_or_404(lesson_id)
    if not lesson:
        flash("Không tìm thấy bài ngữ pháp.")
        return redirect(url_for("grammar_lessons"))

    attempt = get_learning_attempt("grammar_attempts", user["id"], lesson_id) if user["role"] == "student" else None
    if request.method == "POST":
        if user["role"] != "student":
            flash("Chỉ học sinh mới nộp bài ngữ pháp.")
            return redirect(url_for("grammar_detail", lesson_id=lesson_id))

        selected_answer = request.form.get("answer", "").strip().upper()
        if selected_answer not in {"A", "B", "C", "D"}:
            flash("Vui lòng chọn một đáp án.")
            return redirect(url_for("grammar_detail", lesson_id=lesson_id))

        score = 100 if selected_answer == lesson["correct_answer"] else 0
        get_db().execute(
            """
            INSERT INTO grammar_attempts (user_id, lesson_id, selected_answer, score, completed_at)
            VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id, lesson_id)
            DO UPDATE SET selected_answer = excluded.selected_answer,
                          score = excluded.score,
                          completed_at = CURRENT_TIMESTAMP
            """,
            (user["id"], lesson_id, selected_answer, score),
        )
        get_db().commit()
        flash("Đã nộp bài ngữ pháp.")
        attempt = get_learning_attempt("grammar_attempts", user["id"], lesson_id)

    return render_template("grammar_detail.html", user=user, lesson=lesson, attempt=attempt)


@app.route("/skills/grammar/<int:lesson_id>/edit", methods=["GET", "POST"])
def grammar_edit(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chỉ giáo viên mới được chỉnh sửa bài ngữ pháp.")
        return redirect(url_for("grammar_lessons"))

    lesson = get_grammar_lesson_or_404(lesson_id)
    if not lesson:
        flash("Không tìm thấy bài ngữ pháp.")
        return redirect(url_for("grammar_lessons"))

    if request.method == "POST":
        title = request.form.get("title", "").strip()
        content = request.form.get("content", "").strip()
        material_type = request.form.get("material_type", "file").strip()
        material_url = request.form.get("material_url", "").strip()
        material_file = request.files.get("material_file")
        quiz_question = request.form.get("quiz_question", "").strip()
        option_a = request.form.get("option_a", "").strip()
        option_b = request.form.get("option_b", "").strip()
        option_c = request.form.get("option_c", "").strip()
        option_d = request.form.get("option_d", "").strip()
        correct_answer = request.form.get("correct_answer", "").strip().upper()

        if not all([title, content, quiz_question, option_a, option_b, option_c, option_d]):
            flash("Vui lòng nhập đủ tiêu đề, nội dung và câu hỏi ABCD.")
            return redirect(url_for("grammar_edit", lesson_id=lesson_id))
        if correct_answer not in {"A", "B", "C", "D"}:
            flash("Vui lòng tick đáp án đúng.")
            return redirect(url_for("grammar_edit", lesson_id=lesson_id))

        material_path = lesson["material_path"] if material_type == "file" else None
        material_name = lesson["material_name"] if material_type == "file" else None
        if material_type == "link":
            if not material_url:
                flash("Vui lòng nhập link tài liệu.")
                return redirect(url_for("grammar_edit", lesson_id=lesson_id))
            material_url = normalize_embed_url(material_url)
        elif material_type == "file":
            if material_file and material_file.filename:
                if not material_file_is_allowed(material_file.filename):
                    flash("File tài liệu chưa đúng định dạng cho phép.")
                    return redirect(url_for("grammar_edit", lesson_id=lesson_id))
                original_name = secure_filename(material_file.filename)
                saved_name = f"grammar_file_{user['id']}_{int(time.time())}_{original_name}"
                material_file.save(SKILL_FILE_DIR / saved_name)
                material_path = f"uploads/skills/files/{saved_name}"
                material_name = material_file.filename
            elif not material_path:
                flash("Vui lòng upload file tài liệu.")
                return redirect(url_for("grammar_edit", lesson_id=lesson_id))
            material_url = ""
        else:
            flash("Loại tài liệu không hợp lệ.")
            return redirect(url_for("grammar_edit", lesson_id=lesson_id))

        get_db().execute(
            """
            UPDATE grammar_lessons
            SET title = ?, content = ?, material_type = ?, material_url = ?, material_path = ?,
                material_name = ?, quiz_question = ?, option_a = ?, option_b = ?, option_c = ?,
                option_d = ?, correct_answer = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (
                title, content, material_type, material_url or None, material_path, material_name,
                quiz_question, option_a, option_b, option_c, option_d, correct_answer, lesson_id,
            ),
        )
        get_db().commit()
        flash("Đã cập nhật bài ngữ pháp.")
        return redirect(url_for("grammar_lessons"))

    return render_template(
        "grammar_teacher.html",
        user=user,
        lessons=get_all_grammar_lessons(user),
        edit_lesson=lesson,
    )


@app.route("/skills/grammar/<int:lesson_id>/delete", methods=["POST"])
def grammar_delete(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chỉ giáo viên mới được xóa bài ngữ pháp.")
        return redirect(url_for("grammar_lessons"))

    get_db().execute("DELETE FROM grammar_attempts WHERE lesson_id = ?", (lesson_id,))
    get_db().execute("DELETE FROM grammar_lessons WHERE id = ?", (lesson_id,))
    get_db().commit()
    flash("Đã xóa bài ngữ pháp.")
    return redirect(url_for("grammar_lessons"))


@app.route("/skills/writing", methods=["GET", "POST"])
def writing_tasks():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if request.method == "POST":
        if user["role"] != "teacher":
            flash("Chi giao vien moi duoc tao chu de viet.")
            return redirect(url_for("writing_tasks"))

        title = request.form.get("title", "").strip()
        prompt_text = request.form.get("prompt", "").strip()
        try:
            time_limit_minutes = int(request.form.get("time_limit_minutes", "30"))
        except ValueError:
            time_limit_minutes = 30
        time_limit_minutes = max(5, min(180, time_limit_minutes))

        if not title or not prompt_text:
            flash("Vui long nhap day du tieu de va de bai viet.")
            return redirect(url_for("writing_tasks"))

        get_db().execute(
            """
            INSERT INTO writing_tasks (teacher_id, title, prompt, time_limit_minutes)
            VALUES (?, ?, ?, ?)
            """,
            (user["id"], title, prompt_text, time_limit_minutes),
        )
        get_db().commit()
        flash("Da tao chu de viet.")
        return redirect(url_for("writing_tasks"))

    template_name = "writing_teacher.html" if user["role"] == "teacher" else "writing_student.html"
    return render_template(
        template_name,
        user=user,
        tasks=get_all_writing_tasks(user),
        submissions=get_teacher_writing_submissions() if user["role"] == "teacher" else [],
        edit_task=None,
    )


@app.route("/skills/writing/<int:task_id>", methods=["GET", "POST"])
def writing_detail(task_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    task = get_writing_task_or_404(task_id)
    if not task:
        flash("Khong tim thay chu de viet.")
        return redirect(url_for("writing_tasks"))

    if user["role"] != "student":
        return redirect(url_for("writing_tasks"))

    submission = get_writing_submission(user["id"], task_id)
    if request.method == "POST":
        content = request.form.get("content", "").strip()
        if not content:
            flash("Vui long viet bai truoc khi nop.")
            return redirect(url_for("writing_detail", task_id=task_id))

        ai_score, ai_feedback = evaluate_writing_with_ai(task["prompt"], content)
        get_db().execute(
            """
            INSERT INTO writing_submissions
                (task_id, user_id, content, ai_score, ai_feedback, teacher_score,
                 teacher_feedback, status, submitted_at)
            VALUES (?, ?, ?, ?, ?, NULL, NULL, 'waiting_teacher', CURRENT_TIMESTAMP)
            ON CONFLICT(user_id, task_id)
            DO UPDATE SET content = excluded.content,
                          ai_score = excluded.ai_score,
                          ai_feedback = excluded.ai_feedback,
                          teacher_score = NULL,
                          teacher_feedback = NULL,
                          status = 'waiting_teacher',
                          submitted_at = CURRENT_TIMESTAMP,
                          graded_at = NULL
            """,
            (task_id, user["id"], content, ai_score, ai_feedback),
        )
        get_db().commit()
        flash("Da nop bai viet. AI da cham truoc, giao vien se cham sau.")
        submission = get_writing_submission(user["id"], task_id)

    return render_template("writing_detail.html", user=user, task=task, submission=submission)


@app.route("/skills/writing/<int:task_id>/edit", methods=["GET", "POST"])
def writing_edit(task_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc chinh sua chu de viet.")
        return redirect(url_for("writing_tasks"))

    task = get_writing_task_or_404(task_id)
    if not task:
        flash("Khong tim thay chu de viet.")
        return redirect(url_for("writing_tasks"))

    if request.method == "POST":
        title = request.form.get("title", "").strip()
        prompt_text = request.form.get("prompt", "").strip()
        try:
            time_limit_minutes = int(request.form.get("time_limit_minutes", "30"))
        except ValueError:
            time_limit_minutes = 30
        time_limit_minutes = max(5, min(180, time_limit_minutes))

        if not title or not prompt_text:
            flash("Vui long nhap day du tieu de va de bai viet.")
            return redirect(url_for("writing_edit", task_id=task_id))

        get_db().execute(
            """
            UPDATE writing_tasks
            SET title = ?, prompt = ?, time_limit_minutes = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (title, prompt_text, time_limit_minutes, task_id),
        )
        get_db().commit()
        flash("Da cap nhat chu de viet.")
        return redirect(url_for("writing_tasks"))

    return render_template(
        "writing_teacher.html",
        user=user,
        tasks=get_all_writing_tasks(user),
        submissions=get_teacher_writing_submissions(),
        edit_task=task,
    )


@app.route("/skills/writing/<int:task_id>/delete", methods=["POST"])
def writing_delete(task_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc xoa chu de viet.")
        return redirect(url_for("writing_tasks"))

    get_db().execute("DELETE FROM writing_submissions WHERE task_id = ?", (task_id,))
    get_db().execute("DELETE FROM writing_tasks WHERE id = ?", (task_id,))
    get_db().commit()
    flash("Da xoa chu de viet.")
    return redirect(url_for("writing_tasks"))


@app.route("/skills/writing/submissions/<int:submission_id>/grade", methods=["POST"])
def writing_grade(submission_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc cham bai viet.")
        return redirect(url_for("writing_tasks"))

    try:
        teacher_score = int(request.form.get("teacher_score", "0"))
    except ValueError:
        teacher_score = 0
    teacher_score = max(0, min(100, teacher_score))
    teacher_feedback = request.form.get("teacher_feedback", "").strip()

    get_db().execute(
        """
        UPDATE writing_submissions
        SET teacher_score = ?, teacher_feedback = ?, status = 'graded', graded_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """,
        (teacher_score, teacher_feedback, submission_id),
    )
    get_db().commit()
    flash("Da luu diem giao vien.")
    return redirect(url_for("writing_tasks"))


@app.route("/skills/speaking", methods=["GET", "POST"])
def speaking_lessons():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if request.method == "POST":
        if user["role"] != "teacher":
            flash("Chi giao vien moi duoc tao bai noi.")
            return redirect(url_for("speaking_lessons"))

        title = request.form.get("title", "").strip()
        passage = request.form.get("passage", "").strip()
        if not title or not passage:
            flash("Vui long nhap tieu de va doan van noi mau.")
            return redirect(url_for("speaking_lessons"))

        word_notes = generate_word_notes(passage)
        get_db().execute(
            """
            INSERT INTO speaking_lessons (teacher_id, title, passage, word_notes_json)
            VALUES (?, ?, ?, ?)
            """,
            (user["id"], title, passage, json.dumps(word_notes, ensure_ascii=False)),
        )
        get_db().commit()
        flash("Da tao bai noi.")
        return redirect(url_for("speaking_lessons"))

    template_name = "speaking_teacher.html" if user["role"] == "teacher" else "speaking_student.html"
    return render_template(
        template_name,
        user=user,
        lessons=get_all_speaking_lessons(user),
        edit_lesson=None,
    )


@app.route("/skills/speaking/<int:lesson_id>", methods=["GET", "POST"])
def speaking_detail(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    lesson = get_speaking_lesson_or_404(lesson_id)
    if not lesson:
        flash("Khong tim thay bai noi.")
        return redirect(url_for("speaking_lessons"))

    attempt = get_learning_attempt("speaking_attempts", user["id"], lesson_id) if user["role"] == "student" else None
    word_notes = get_word_notes(lesson)
    if request.method == "POST":
        if user["role"] != "student":
            flash("Chi hoc sinh moi duoc nop bai noi.")
            return redirect(url_for("speaking_detail", lesson_id=lesson_id))

        transcript = request.form.get("transcript", "").strip()
        try:
            score = int(float(request.form.get("score", "0")))
        except ValueError:
            score = 0
        score = max(0, min(100, score))
        missed_words_json = request.form.get("missed_words_json", "[]")
        try:
            missed_words = json.loads(missed_words_json)
            if not isinstance(missed_words, list):
                missed_words = []
        except json.JSONDecodeError:
            missed_words = []

        get_db().execute(
            """
            INSERT INTO speaking_attempts
                (user_id, lesson_id, transcript, score, missed_words_json, completed_at)
            VALUES (?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(user_id, lesson_id)
            DO UPDATE SET transcript = excluded.transcript,
                          score = excluded.score,
                          missed_words_json = excluded.missed_words_json,
                          completed_at = CURRENT_TIMESTAMP
            """,
            (user["id"], lesson_id, transcript, score, json.dumps(missed_words, ensure_ascii=False)),
        )
        get_db().commit()
        flash("Da luu ket qua bai noi.")
        attempt = get_learning_attempt("speaking_attempts", user["id"], lesson_id)

    attempt_missed_words = []
    if attempt:
        try:
            attempt_missed_words = json.loads(attempt["missed_words_json"] or "[]")
            if not isinstance(attempt_missed_words, list):
                attempt_missed_words = []
        except json.JSONDecodeError:
            attempt_missed_words = []

    return render_template(
        "speaking_detail.html",
        user=user,
        lesson=lesson,
        attempt=attempt,
        word_notes=word_notes,
        attempt_missed_words=attempt_missed_words,
        expected_words=normalize_practice_words(lesson["passage"]),
    )


@app.route("/skills/speaking/<int:lesson_id>/edit", methods=["GET", "POST"])
def speaking_edit(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc chinh sua bai noi.")
        return redirect(url_for("speaking_lessons"))

    lesson = get_speaking_lesson_or_404(lesson_id)
    if not lesson:
        flash("Khong tim thay bai noi.")
        return redirect(url_for("speaking_lessons"))

    if request.method == "POST":
        title = request.form.get("title", "").strip()
        passage = request.form.get("passage", "").strip()
        if not title or not passage:
            flash("Vui long nhap tieu de va doan van noi mau.")
            return redirect(url_for("speaking_edit", lesson_id=lesson_id))

        word_notes = generate_word_notes(passage)
        get_db().execute(
            """
            UPDATE speaking_lessons
            SET title = ?, passage = ?, word_notes_json = ?, updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            (title, passage, json.dumps(word_notes, ensure_ascii=False), lesson_id),
        )
        get_db().commit()
        flash("Da cap nhat bai noi.")
        return redirect(url_for("speaking_lessons"))

    return render_template(
        "speaking_teacher.html",
        user=user,
        lessons=get_all_speaking_lessons(user),
        edit_lesson=lesson,
    )


@app.route("/skills/speaking/<int:lesson_id>/delete", methods=["POST"])
def speaking_delete(lesson_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chi giao vien moi duoc xoa bai noi.")
        return redirect(url_for("speaking_lessons"))

    get_db().execute("DELETE FROM speaking_attempts WHERE lesson_id = ?", (lesson_id,))
    get_db().execute("DELETE FROM speaking_lessons WHERE id = ?", (lesson_id,))
    get_db().commit()
    flash("Da xoa bai noi.")
    return redirect(url_for("speaking_lessons"))


@app.route("/games/quizz", methods=["GET", "POST"])
def quizz():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if request.method == "POST":
        if user["role"] != "teacher":
            flash("Chỉ giáo viên mới được tạo gói quiz.")
            return redirect(url_for("quizz"))

        title = request.form.get("title", "").strip()
        difficulty = request.form.get("difficulty", "").strip()
        questions = get_quiz_form_questions()

        if not title or not difficulty:
            flash("Vui lòng nhập tên gói và độ khó.")
            return redirect(url_for("quizz"))
        if questions is None:
            flash("Một gói quiz phải có đủ 10 câu, mỗi câu đủ 4 đáp án và đáp án đúng.")
            return redirect(url_for("quizz"))

        db = get_db()
        cursor = db.execute(
            "INSERT INTO quiz_packages (teacher_id, title, difficulty) VALUES (?, ?, ?)",
            (user["id"], title, difficulty),
        )
        save_quiz_questions(cursor.lastrowid, questions)
        db.commit()
        flash("Đã tạo gói quiz mới.")
        return redirect(url_for("quizz"))

    template_name = "quizz_teacher.html" if user["role"] == "teacher" else "quizz_student.html"
    return render_template(template_name, user=user, packages=get_all_quiz_packages(user), edit_package=None, edit_questions=[])


@app.route("/games/rank-kind")
def rank_kind():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "student":
        flash("Rank kind dành cho học sinh.")
        return redirect(url_for("quizz"))
    return render_template(
        "rank_kind.html",
        user=user,
        packages=get_rank_kind_packages(user),
        leaderboard=get_rank_kind_leaderboard(),
    )


@app.get("/api/rank-kind/leaderboard")
def rank_kind_leaderboard_api():
    user = get_current_user()
    if not user:
        return jsonify({"ok": False, "message": "Chưa đăng nhập."}), 401
    return jsonify({"ok": True, "leaderboard": get_rank_kind_leaderboard()})


@app.post("/api/rank-kind/queue/join")
def rank_kind_queue_join():
    user = get_current_user()
    if not user:
        return jsonify({"ok": False, "message": "Chưa đăng nhập."}), 401
    if user["role"] != "student":
        return jsonify({"ok": False, "message": "Chỉ học sinh mới được ghép trận."}), 403

    data = request.get_json(silent=True) or {}
    package_id = int(data.get("package_id") or 0)
    package = get_rank_kind_package_or_first(package_id, user)
    if not package:
        return jsonify({"ok": False, "message": "Chưa có gói quiz đủ 10 câu."}), 400

    player = get_rank_kind_player(user)
    with RANK_KIND_LOCK:
        RANK_KIND_QUEUE[:] = [ticket for ticket in RANK_KIND_QUEUE if ticket["player"]["id"] != player["id"]]
        for match in RANK_KIND_MATCHES.values():
            if match["state"] != "finished" and any(item["id"] == player["id"] for item in match["players"]):
                return jsonify({"ok": True, **serialize_rank_kind_match(match, user["id"])})

        opponent_index = next(
            (
                index
                for index, ticket in enumerate(RANK_KIND_QUEUE)
                if ticket["player"]["id"] != player["id"] and ticket["packageId"] == package["id"]
            ),
            None,
        )
        if opponent_index is not None:
            ticket = RANK_KIND_QUEUE.pop(opponent_index)
            match = create_rank_kind_match(ticket, player, package)
            return jsonify({"ok": True, **serialize_rank_kind_match(match, user["id"])})

        RANK_KIND_QUEUE.append(
            {
                "player": player,
                "packageId": package["id"],
                "packageTitle": package["title"],
                "joinedAt": time.time(),
            }
        )
        return jsonify(
            {
                "ok": True,
                "status": "waiting",
                "packageTitle": package["title"],
                "localPlayerId": player["id"],
                "joinedAt": time.time(),
            }
        )


@app.post("/api/rank-kind/queue/leave")
def rank_kind_queue_leave():
    user = get_current_user()
    if not user:
        return jsonify({"ok": False, "message": "Chưa đăng nhập."}), 401
    player_id = str(user["id"])
    with RANK_KIND_LOCK:
        RANK_KIND_QUEUE[:] = [ticket for ticket in RANK_KIND_QUEUE if ticket["player"]["id"] != player_id]
        for match_id, match in list(RANK_KIND_MATCHES.items()):
            if not any(player["id"] == player_id for player in match["players"]):
                continue
            if match["state"] == "ready":
                RANK_KIND_MATCHES.pop(match_id, None)
            elif match["state"] == "playing":
                opponent = next((player for player in match["players"] if player["id"] != player_id), None)
                if opponent:
                    match["winnerId"] = opponent["id"]
                    finish_rank_kind_match(match)
    return jsonify({"ok": True, "status": "left"})


@app.post("/api/rank-kind/ready")
def rank_kind_ready():
    user = get_current_user()
    if not user:
        return jsonify({"ok": False, "message": "Chưa đăng nhập."}), 401
    data = request.get_json(silent=True) or {}
    match_id = data.get("match_id", "")
    with RANK_KIND_LOCK:
        match = RANK_KIND_MATCHES.get(match_id)
        if not match:
            return jsonify({"ok": False, "message": "Không tìm thấy trận."}), 404
        if not any(player["id"] == str(user["id"]) for player in match["players"]):
            return jsonify({"ok": False, "message": "Bạn không ở trong trận này."}), 403
        match["ready"].add(str(user["id"]))
        if match["state"] == "ready" and len(match["ready"]) >= 2:
            match["state"] = "playing"
            start_rank_kind_question(match, 0)
        return jsonify({"ok": True, **serialize_rank_kind_match(match, user["id"])})


@app.get("/api/rank-kind/state")
def rank_kind_state():
    user = get_current_user()
    if not user:
        return jsonify({"ok": False, "message": "Chưa đăng nhập."}), 401
    match_id = request.args.get("match_id", "")
    player_id = str(user["id"])
    with RANK_KIND_LOCK:
        if match_id and match_id in RANK_KIND_MATCHES:
            return jsonify({"ok": True, **serialize_rank_kind_match(RANK_KIND_MATCHES[match_id], user["id"])})
        for match in RANK_KIND_MATCHES.values():
            if match["state"] != "finished" and any(player["id"] == player_id for player in match["players"]):
                return jsonify({"ok": True, **serialize_rank_kind_match(match, user["id"])})
        ticket = next((item for item in RANK_KIND_QUEUE if item["player"]["id"] == player_id), None)
        if ticket:
            return jsonify(
                {
                    "ok": True,
                    "status": "waiting",
                    "packageTitle": ticket["packageTitle"],
                    "localPlayerId": player_id,
                    "joinedAt": ticket["joinedAt"],
                }
            )
    return jsonify({"ok": True, "status": "idle", "leaderboard": get_rank_kind_leaderboard()})


@app.post("/api/rank-kind/answer")
def rank_kind_answer():
    user = get_current_user()
    if not user:
        return jsonify({"ok": False, "message": "Chưa đăng nhập."}), 401
    data = request.get_json(silent=True) or {}
    match_id = data.get("match_id", "")
    question_id = str(data.get("question_id", ""))
    selected_key = str(data.get("selected_key", "")).upper()
    player_id = str(user["id"])
    with RANK_KIND_LOCK:
        match = RANK_KIND_MATCHES.get(match_id)
        if not match:
            return jsonify({"ok": False, "message": "Không tìm thấy trận."}), 404
        if match["state"] != "playing" or not match.get("current"):
            return jsonify({"ok": True, **serialize_rank_kind_match(match, user["id"])})
        current = match["current"]
        if current["id"] != question_id or current.get("resolved"):
            return jsonify({"ok": True, **serialize_rank_kind_match(match, user["id"])})
        if player_id in current["answered"]:
            return jsonify({"ok": True, **serialize_rank_kind_match(match, user["id"])})

        current["answered"].add(player_id)
        question = match["questions"][match["questionIndex"]]
        if selected_key == question["correctKey"]:
            resolve_rank_kind_question(
                match,
                winner_id=player_id,
                selected_by=player_id,
                selected_key=selected_key,
                reason="correct_fastest",
            )
        elif len(current["answered"]) >= 2:
            resolve_rank_kind_question(
                match,
                selected_by=player_id,
                selected_key=selected_key,
                reason="both_wrong",
            )
        return jsonify({"ok": True, **serialize_rank_kind_match(match, user["id"])})


@app.route("/games/quizz/<int:package_id>/edit", methods=["GET", "POST"])
def quizz_edit(package_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chỉ giáo viên mới được chỉnh sửa gói quiz.")
        return redirect(url_for("quizz"))

    package = get_quiz_package_or_404(package_id)
    if not package:
        flash("Không tìm thấy gói quiz.")
        return redirect(url_for("quizz"))

    if request.method == "POST":
        title = request.form.get("title", "").strip()
        difficulty = request.form.get("difficulty", "").strip()
        questions = get_quiz_form_questions()

        if not title or not difficulty:
            flash("Vui lòng nhập tên gói và độ khó.")
            return redirect(url_for("quizz_edit", package_id=package_id))
        if questions is None:
            flash("Một gói quiz phải có đủ 10 câu, mỗi câu đủ 4 đáp án và đáp án đúng.")
            return redirect(url_for("quizz_edit", package_id=package_id))

        db = get_db()
        db.execute(
            "UPDATE quiz_packages SET title = ?, difficulty = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            (title, difficulty, package_id),
        )
        save_quiz_questions(package_id, questions)
        db.commit()
        flash("Đã cập nhật gói quiz.")
        return redirect(url_for("quizz"))

    return render_template(
        "quizz_teacher.html",
        user=user,
        packages=get_all_quiz_packages(user),
        edit_package=package,
        edit_questions=get_quiz_questions(package_id),
    )


@app.route("/games/quizz/<int:package_id>/delete", methods=["POST"])
def quizz_delete(package_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))
    if user["role"] != "teacher":
        flash("Chỉ giáo viên mới được xóa gói quiz.")
        return redirect(url_for("quizz"))

    package = get_quiz_package_or_404(package_id)
    if not package:
        flash("Không tìm thấy gói quiz.")
        return redirect(url_for("quizz"))

    db = get_db()
    db.execute("DELETE FROM quiz_attempts WHERE package_id = ?", (package_id,))
    db.execute("DELETE FROM quiz_questions WHERE package_id = ?", (package_id,))
    db.execute("DELETE FROM quiz_packages WHERE id = ?", (package_id,))
    db.commit()
    flash("Đã xóa gói quiz.")
    return redirect(url_for("quizz"))


@app.route("/games/quizz/<int:package_id>/play", methods=["GET", "POST"])
def quizz_play(package_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    package = get_quiz_package_or_404(package_id)
    if not package:
        flash("Không tìm thấy gói quiz.")
        return redirect(url_for("quizz"))

    if request.method == "POST":
        question_ids = [int(item) for item in request.form.get("question_ids", "").split(",") if item.isdigit()]
        if len(question_ids) != 10:
            flash("Lượt chơi quiz chưa hợp lệ.")
            return redirect(url_for("quizz_play", package_id=package_id))

        placeholders = ",".join("?" for _ in question_ids)
        rows = get_db().execute(
            f"SELECT id, correct_answer FROM quiz_questions WHERE package_id = ? AND id IN ({placeholders})",
            (package_id, *question_ids),
        ).fetchall()
        correct_map = {row["id"]: row["correct_answer"] for row in rows}
        correct_count = 0
        for question_id in question_ids:
            selected_answer = request.form.get(f"answer_{question_id}", "").strip().upper()
            if selected_answer == correct_map.get(question_id):
                correct_count += 1

        total_questions = len(question_ids)
        percentage = round(correct_count / total_questions * 100)
        bonus_points = 10 if percentage > 80 else 0
        get_db().execute(
            """
            INSERT INTO quiz_attempts
                (user_id, package_id, correct_count, total_questions, percentage, bonus_points)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (user["id"], package_id, correct_count, total_questions, percentage, bonus_points),
        )
        get_db().commit()
        flash("Bạn được cộng 10 điểm bảng xếp hạng." if bonus_points else "Bạn cần đúng trên 80% để được cộng điểm.")
        return render_template(
            "quizz_result.html",
            user=user,
            package=package,
            correct_count=correct_count,
            total_questions=total_questions,
            percentage=percentage,
            bonus_points=bonus_points,
        )

    questions = list(get_quiz_questions(package_id))
    if len(questions) != 10:
        flash("Gói quiz này chưa đủ 10 câu.")
        return redirect(url_for("quizz"))
    random.shuffle(questions)
    return render_template("quizz_play.html", user=user, package=package, questions=questions)


@app.route("/games/frog")
def frog_game():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    packages = get_all_quiz_packages(user)
    package = next((item for item in packages if item["question_count"] == 10), None)
    if not package:
        flash("Chưa có gói quiz đủ 10 câu để chơi Ếch vui vẻ.")
        return redirect(url_for("quizz"))
    return redirect(url_for("frog_game_play", package_id=package["id"]))


@app.route("/games/frog/<int:package_id>/play")
def frog_game_play(package_id):
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    package = get_quiz_package_or_404(package_id)
    if not package:
        flash("Không tìm thấy gói quiz.")
        return redirect(url_for("quizz"))

    rows = list(get_quiz_questions(package_id))
    if len(rows) != 10:
        flash("Gói quiz này chưa đủ 10 câu.")
        return redirect(url_for("quizz"))

    questions = []
    for index, row in enumerate(rows, start=1):
        options = [
            {"key": "A", "text": row["option_a"]},
            {"key": "B", "text": row["option_b"]},
            {"key": "C", "text": row["option_c"]},
            {"key": "D", "text": row["option_d"]},
        ]
        questions.append(
            {
                "number": index,
                "text": row["question_text"],
                "options": options,
                "correct": row["correct_answer"],
            }
        )

    return render_template("frog_game.html", user=user, package=package, questions=questions)


@app.route("/profile", methods=["GET", "POST"])
def profile():
    user = get_current_user()
    if not user:
        return redirect(url_for("login"))

    if request.method == "POST":
        fullname = request.form.get("fullname", "").strip()
        username = request.form.get("username", "").strip() or None
        email = request.form.get("email", "").strip().lower()
        phone = request.form.get("phone", "").strip() or None
        class_name = request.form.get("class", "").strip() or None

        if not fullname or not email:
            flash("Vui lòng điền họ tên và email.")
            return redirect(url_for("profile"))

        db = get_db()
        duplicate = db.execute(
            "SELECT id FROM users WHERE id != ? AND (email = ? OR username = ?)",
            (user["id"], email, username),
        ).fetchone()
        if duplicate:
            flash("Email hoặc tên đăng nhập đã được sử dụng.")
            return redirect(url_for("profile"))

        avatar_path = user["avatar_path"]
        avatar_file = request.files.get("avatar")
        if avatar_file and avatar_file.filename:
            if not avatar_is_allowed(avatar_file.filename):
                flash("Ảnh đại diện chỉ nhận PNG, JPG, JPEG, WEBP hoặc GIF.")
                return redirect(url_for("profile"))

            filename = secure_filename(avatar_file.filename)
            extension = filename.rsplit(".", 1)[1].lower()
            saved_name = f"user_{user['id']}_avatar_{int(time.time())}.{extension}"
            save_path = AVATAR_DIR / saved_name
            avatar_file.save(save_path)
            avatar_path = f"uploads/avatars/{saved_name}"

        db.execute(
            """
            UPDATE users
            SET fullname = ?, username = ?, email = ?, phone = ?, class = ?, avatar_path = ?
            WHERE id = ?
            """,
            (fullname, username, email, phone, class_name, avatar_path, user["id"]),
        )
        db.commit()
        session["username"] = username or email
        flash("Đã lưu thông tin hồ sơ.")
        return redirect(url_for("profile"))

    return render_template("profile_student.html", user=user)


@app.route("/dashboard")
def dashboard():
    return redirect(url_for("home"))


if __name__ == "__main__":
    with app.app_context():
        init_db()
    app.run(
        host=os.environ.get("HOST", "127.0.0.1"),
        port=int(os.environ.get("PORT", 5000)),
        debug=os.environ.get("FLASK_DEBUG") == "1",
    )
